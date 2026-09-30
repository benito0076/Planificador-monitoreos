#!/usr/bin/env python3
"""
Planificador de Monitoreos Ambientales
======================================

Herramienta de consola para programar campañas de monitoreo ambiental:

  * Personal: competencias por matriz, conductores, días no disponibles,
    vencimiento de habilitaciones (EMO / SCTR / inducción) y descanso obligatorio.
  * Vehículos: capacidad, 4x4, velocidad, rendimiento, costo diario,
    mantenimiento y vencimiento de documentos (SOAT / revisión técnica).
  * Equipos: inventario por unidad con vencimiento de calibración; cada punto
    indica qué tipos de equipo necesita y se asignan a las cuadrillas por día.
  * Puntos: coordenadas, matriz, duración, acceso a pie, prioridad, ventana
    horaria, tiempo máximo de preservación de muestras, frecuencia de monitoreo
    y muestreos de larga exposición (p. ej. PM10 24 h: instalación + retiro).
  * Tiempos reales por carretera vía OSRM (con caché); sin conexión se estiman
    con Haversine x factor de ruta.
  * Formación de cuadrillas balanceando la carga de trabajo, ruteo diario
    (vecino más cercano + 2-opt) respetando jornada, almuerzo, ventanas,
    preservación y disponibilidad de equipos.
  * Costos (combustible, vehículo, viáticos), alertas de vencimientos,
    historial de programaciones.
  * Salidas: CSV, Excel (openpyxl), hoja de ruta imprimible HTML con enlaces
    de navegación y mapa HTML con rutas por carretera (folium).

Uso:
    python planificador_monitoreo.py          # menú interactivo
    python planificador_monitoreo.py --demo   # carga datos de ejemplo y genera un plan

Importación masiva: en cada módulo use "Exportar CSV" para obtener la plantilla
con todas las columnas, edítela en Excel y vuelva a importarla. Las listas
(competencias, equipos, fechas) se separan con "|".
"""

from __future__ import annotations

import calendar
import csv
import html
import json
import math
import sys
import time as time_mod
import urllib.error
import urllib.request
from dataclasses import MISSING, asdict, dataclass, field, fields
from datetime import date, datetime, time, timedelta
from pathlib import Path

# En un .exe (PyInstaller) __file__ apunta a una carpeta temporal: guardar junto al ejecutable.
CARPETA = Path(sys.executable).parent if getattr(sys, "frozen", False) else Path(__file__).parent
ARCHIVO_DATOS = CARPETA / "datos_monitoreo.json"
ARCHIVO_CACHE_OSRM = CARPETA / "cache_osrm.json"
CARPETA_SALIDAS = CARPETA / "salidas"
CARPETA_HISTORIAL = CARPETA / "historial"
UMBRAL_SNAP_M = 500  # avisar si un punto queda a más de esto de la vía más cercana
DIAS_ALERTA = 30     # anticipación de las alertas de vencimiento
MATRICES = ["agua", "aire", "suelo", "ruido", "sedimento", "biologico", "efluente"]
DIAS_SEMANA = ["Lun", "Mar", "Mié", "Jue", "Vie", "Sáb", "Dom"]
# Multiplicador de "costo" para elegir el siguiente punto: menor = se atiende antes.
# 0 = retiro de equipos instalados (no puede quedar pendiente).
PESO_PRIORIDAD = {0: 0.3, 1: 0.5, 2: 1.0, 3: 1.5}
# Frecuencia -> (meses, días) entre monitoreos; None = monitoreo único.
FRECUENCIAS = {"unica": None, "semanal": (0, 7), "quincenal": (0, 15), "mensual": (1, 0),
               "bimestral": (2, 0), "trimestral": (3, 0), "semestral": (6, 0), "anual": (12, 0)}
SI = {"si", "sí", "s", "1", "true", "x", "yes", "y"}
BASE = "__BASE__"


def vigente(hasta: str, dia: date) -> bool:
    return not hasta or date.fromisoformat(hasta) >= dia


def sumar_meses(d: date, n: int) -> date:
    m = d.month - 1 + n
    y, m = d.year + m // 12, m % 12 + 1
    return date(y, m, min(d.day, calendar.monthrange(y, m)[1]))


# --------------------------------------------------------------------------- #
# Modelo de datos
# --------------------------------------------------------------------------- #
@dataclass
class Persona:
    nombre: str
    cargo: str = "Técnico"
    competencias: list[str] = field(default_factory=list)
    conduce: bool = False
    activo: bool = True
    habilitado_hasta: str = ""  # vencimiento de EMO / SCTR / inducción (AAAA-MM-DD)
    no_disponible: list[str] = field(default_factory=list)  # fechas ISO

    def disponible(self, dia: date) -> bool:
        return self.activo and dia.isoformat() not in self.no_disponible and vigente(self.habilitado_hasta, dia)


@dataclass
class Vehiculo:
    placa: str
    descripcion: str = "Camioneta"
    capacidad: int = 4
    es_4x4: bool = True
    velocidad_kmh: float = 45.0
    rendimiento_km_l: float = 10.0
    costo_dia: float = 0.0          # alquiler / depreciación por día de uso
    activo: bool = True
    documentos_hasta: str = ""      # vencimiento más próximo de SOAT / revisión técnica
    no_disponible: list[str] = field(default_factory=list)  # mantenimiento, etc.

    def disponible(self, dia: date) -> bool:
        return self.activo and dia.isoformat() not in self.no_disponible and vigente(self.documentos_hasta, dia)


@dataclass
class Equipo:
    codigo: str
    tipo: str                       # p. ej. multiparametro, sonometro, muestreador_pm10
    descripcion: str = ""
    calibracion_vence: str = ""
    activo: bool = True
    no_disponible: list[str] = field(default_factory=list)

    def disponible(self, dia: date) -> bool:
        return self.activo and dia.isoformat() not in self.no_disponible and vigente(self.calibracion_vence, dia)


@dataclass
class Punto:
    codigo: str
    nombre: str
    lat: float
    lon: float
    matriz: str = "agua"
    duracion_min: int = 60          # muestreo, o instalación si hay exposición
    requiere_4x4: bool = False
    acceso_min: int = 0             # caminata ida+vuelta desde donde se deja el vehículo
    prioridad: int = 2              # 1 alta, 2 media, 3 baja
    personas_min: int = 2
    activo: bool = True
    hora_desde: str = ""            # ventana horaria para muestrear ("" = sin límite)
    hora_hasta: str = ""
    max_horas_preservacion: float = 0   # horas máx. desde la toma hasta llegar a base (0 = sin límite)
    equipos: list[str] = field(default_factory=list)  # tipos de equipo necesarios
    horas_exposicion: float = 0     # > 0: se instala y se retira después (p. ej. 24 h)
    duracion_retiro_min: int = 30
    frecuencia: str = "unica"
    ultimo_monitoreo: str = ""

    def proxima_fecha(self) -> date | None:
        """Fecha en que corresponde el siguiente monitoreo; date.min = nunca realizado;
        None = monitoreo único ya realizado."""
        if not self.ultimo_monitoreo:
            return date.min
        f = FRECUENCIAS.get(self.frecuencia)
        if f is None:
            return None
        return sumar_meses(date.fromisoformat(self.ultimo_monitoreo), f[0]) + timedelta(days=f[1])


@dataclass
class Configuracion:
    base_nombre: str = "Base de operaciones"
    base_lat: float = -12.0464
    base_lon: float = -77.0428
    hora_inicio: str = "07:00"
    hora_fin: str = "17:00"
    preparacion_min: int = 30      # carga de equipos, calibración, charla de seguridad
    descarga_min: int = 30         # descarga, preservación y registro de muestras al volver
    hora_almuerzo: str = "12:30"
    almuerzo_min: int = 45
    factor_ruta: float = 1.35      # distancia real / línea recta (solo si no hay ruta OSRM)
    usar_osrm: bool = True
    osrm_url: str = "https://router.project-osrm.org"
    factor_tiempo_osrm: float = 1.2  # OSRM asume auto a velocidad legal; ajuste por tráfico/camioneta
    dias_laborables: list[int] = field(default_factory=lambda: [0, 1, 2, 3, 4, 5])
    tamano_cuadrilla: int = 3      # personas por cuadrilla (incluye conductor)
    max_dias_consecutivos: int = 6
    moneda: str = "S/"
    precio_combustible: float = 4.3  # por litro
    viatico_persona_dia: float = 60.0


def _crear(cls, d: dict):
    """Crea el objeto ignorando campos desconocidos (compatibilidad entre versiones)."""
    nombres = {f.name for f in fields(cls)}
    return cls(**{k: v for k, v in d.items() if k in nombres})


class Repositorio:
    def __init__(self, ruta: Path = ARCHIVO_DATOS):
        self.ruta = ruta
        self.personal: list[Persona] = []
        self.vehiculos: list[Vehiculo] = []
        self.equipos: list[Equipo] = []
        self.puntos: list[Punto] = []
        self.config = Configuracion()
        self.cargar()

    def cargar(self) -> None:
        if not self.ruta.exists():
            return
        d = json.loads(self.ruta.read_text(encoding="utf-8"))
        self.personal = [_crear(Persona, x) for x in d.get("personal", [])]
        self.vehiculos = [_crear(Vehiculo, x) for x in d.get("vehiculos", [])]
        self.equipos = [_crear(Equipo, x) for x in d.get("equipos", [])]
        self.puntos = [_crear(Punto, x) for x in d.get("puntos", [])]
        self.config = _crear(Configuracion, d.get("config", {}))

    def guardar(self) -> None:
        d = {
            "personal": [asdict(x) for x in self.personal],
            "vehiculos": [asdict(x) for x in self.vehiculos],
            "equipos": [asdict(x) for x in self.equipos],
            "puntos": [asdict(x) for x in self.puntos],
            "config": asdict(self.config),
        }
        self.ruta.write_text(json.dumps(d, indent=2, ensure_ascii=False), encoding="utf-8")

    def tipos_equipo(self) -> list[str]:
        return sorted({e.tipo for e in self.equipos})


# --------------------------------------------------------------------------- #
# Geografía y tiempos
# --------------------------------------------------------------------------- #
def haversine_km(lat1: float, lon1: float, lat2: float, lon2: float) -> float:
    r = 6371.0
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp, dl = p2 - p1, math.radians(lon2 - lon1)
    a = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * r * math.asin(math.sqrt(a))


def _leer_cache_osrm() -> dict:
    try:
        return json.loads(ARCHIVO_CACHE_OSRM.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def _guardar_cache_osrm(cache: dict) -> None:
    try:
        ARCHIVO_CACHE_OSRM.write_text(json.dumps(cache), encoding="utf-8")
    except OSError:
        pass


class Distancias:
    """Distancias (km) y tiempos (min) entre la base y los puntos.

    Con OSRM activo usa rutas reales por carretera (tabla OSRM, con caché en disco).
    Cualquier par sin ruta —o si el servidor no responde— cae a la estimación
    Haversine x factor de ruta con la velocidad promedio del vehículo."""

    def __init__(self, puntos: list[Punto], cfg: Configuracion, verbose: bool = True):
        self.coord = {p.codigo: (p.lat, p.lon) for p in puntos}
        self.coord[BASE] = (cfg.base_lat, cfg.base_lon)
        self.factor = cfg.factor_ruta
        self.factor_tiempo = cfg.factor_tiempo_osrm
        self._cache: dict[tuple[str, str], float] = {}
        self.rutas: dict[tuple[str, str], tuple[float, float]] = {}  # (a, b) -> (km, min) OSRM
        self.fuente = f"estimación en línea recta x factor de ruta {cfg.factor_ruta}"
        self.avisos: list[str] = []
        if cfg.usar_osrm and len(self.coord) > 1:
            self._cargar_osrm(cfg.osrm_url.rstrip("/"))
        if verbose:
            print(f"\n  Tiempos de traslado: {self.fuente}")
            for a in self.avisos:
                print(f"  ⚠ {a}")

    # --- OSRM ------------------------------------------------------------- #
    def _clave(self, cod: str) -> str:
        lat, lon = self.coord[cod]
        return f"{lat:.5f},{lon:.5f}"

    def _cargar_osrm(self, url: str) -> None:
        cache = _leer_cache_osrm()
        par = lambda a, b: f"{url}|{self._clave(a)}|{self._clave(b)}"
        cods = list(self.coord)
        faltan = [(a, b) for a in cods for b in cods if a != b and par(a, b) not in cache]
        if faltan:
            fuentes = list(dict.fromkeys(a for a, _ in faltan))
            destinos = list(dict.fromkeys(b for _, b in faltan))
            print(f"  Consultando OSRM ({len(faltan)} tramos)...", flush=True)
            try:
                self._consultar_osrm(url, fuentes, destinos, cache, par)
            except (urllib.error.URLError, TimeoutError, OSError, ValueError, KeyError) as e:
                self.avisos.append(f"OSRM no disponible ({e}); tramos sin caché usan tiempo estimado.")
            _guardar_cache_osrm(cache)

        sin_ruta = 0
        for a in cods:
            for b in cods:
                if a == b:
                    continue
                v = cache.get(par(a, b))
                if v and v[0] is not None and v[1] is not None:
                    self.rutas[(a, b)] = (v[0], v[1])
                else:
                    sin_ruta += 1
        for c in cods:
            snap = cache.get(f"{url}|snap|{self._clave(c)}")
            if snap and snap > UMBRAL_SNAP_M:
                nombre = "BASE" if c == BASE else c
                self.avisos.append(f"{nombre} está a {snap / 1000:.2f} km de la vía más cercana; "
                                   "revise coordenadas o el tiempo de acceso a pie.")
        if self.rutas:
            self.fuente = f"OSRM por carretera ({url}) x{self.factor_tiempo}"
            if sin_ruta:
                self.avisos.append(f"{sin_ruta} tramo(s) sin ruta OSRM; se usa la estimación en línea recta.")

    def _consultar_osrm(self, url, fuentes, destinos, cache, par) -> None:
        bloque = 50  # 50 orígenes + 50 destinos <= 100 coordenadas por consulta
        primera = True
        for i in range(0, len(fuentes), bloque):
            for j in range(0, len(destinos), bloque):
                src, dst = fuentes[i:i + bloque], destinos[j:j + bloque]
                cods = list(dict.fromkeys(src + dst))
                idx = {c: k for k, c in enumerate(cods)}
                locs = ";".join(f"{self.coord[c][1]:.6f},{self.coord[c][0]:.6f}" for c in cods)
                q = (f"{url}/table/v1/driving/{locs}?annotations=duration,distance"
                     f"&sources={';'.join(str(idx[c]) for c in src)}"
                     f"&destinations={';'.join(str(idx[c]) for c in dst)}")
                if not primera:
                    time_mod.sleep(1)  # el servidor público admite ~1 consulta/segundo
                primera = False
                req = urllib.request.Request(q, headers={"User-Agent": "planificador-monitoreo/1.0"})
                with urllib.request.urlopen(req, timeout=30) as r:
                    d = json.load(r)
                if d.get("code") != "Ok":
                    raise ValueError(d.get("message") or d.get("code"))
                for si, a in enumerate(src):
                    for di, b in enumerate(dst):
                        if a != b:
                            dis, dur = d["distances"][si][di], d["durations"][si][di]
                            cache[par(a, b)] = [None if dis is None else dis / 1000,
                                                None if dur is None else dur / 60]
                for wp, c in zip(d.get("sources", []), src):
                    cache[f"{url}|snap|{self._clave(c)}"] = wp.get("distance", 0)
                for wp, c in zip(d.get("destinations", []), dst):
                    cache[f"{url}|snap|{self._clave(c)}"] = wp.get("distance", 0)

    # --- Consultas -------------------------------------------------------- #
    def km(self, a: str, b: str) -> float:
        if a == b:
            return 0.0
        if (a, b) in self.rutas:
            return self.rutas[(a, b)][0]
        k = (a, b) if a < b else (b, a)
        if k not in self._cache:
            self._cache[k] = haversine_km(*self.coord[a], *self.coord[b]) * self.factor
        return self._cache[k]

    def minutos(self, a: str, b: str, velocidad_kmh: float) -> float:
        if a == b:
            return 0.0
        if (a, b) in self.rutas:
            return self.rutas[(a, b)][1] * self.factor_tiempo
        return self.km(a, b) / velocidad_kmh * 60

    def km_ruta(self, codigos: list[str]) -> float:
        seq = [BASE, *codigos, BASE]
        return sum(self.km(x, y) for x, y in zip(seq, seq[1:]))


_ultima_consulta_ruta = 0.0


def trazado_osrm(coords: list[tuple[float, float]], url: str, cache: dict) -> list[list[float]] | None:
    """Trazado por carretera [[lat, lon], ...] que pasa por `coords` en orden (OSRM /route).
    Devuelve None si no hay ruta o el servidor no responde."""
    url = url.rstrip("/")
    clave = f"{url}|route|" + ";".join(f"{lat:.5f},{lon:.5f}" for lat, lon in coords)
    if clave in cache:
        return cache[clave]
    locs = ";".join(f"{lon:.6f},{lat:.6f}" for lat, lon in coords)
    q = f"{url}/route/v1/driving/{locs}?overview=full&geometries=geojson"
    global _ultima_consulta_ruta
    espera = 1.0 - (time_mod.monotonic() - _ultima_consulta_ruta)
    if espera > 0:
        time_mod.sleep(espera)  # servidor público: ~1 consulta/segundo
    _ultima_consulta_ruta = time_mod.monotonic()
    try:
        req = urllib.request.Request(q, headers={"User-Agent": "planificador-monitoreo/1.0"})
        with urllib.request.urlopen(req, timeout=30) as r:
            d = json.load(r)
        if d.get("code") != "Ok" or not d.get("routes"):
            return None
        # GeoJSON viene como [lon, lat]; se redondea a ~1 m para aligerar caché y HTML.
        linea = [[round(lat, 5), round(lon, 5)] for lon, lat in d["routes"][0]["geometry"]["coordinates"]]
    except (urllib.error.URLError, TimeoutError, OSError, ValueError, KeyError):
        return None
    cache[clave] = linea
    return linea


def parse_hora(s: str) -> time:
    h, m = s.strip().split(":")
    return time(int(h), int(m))


def fmt_min(minutos: float) -> str:
    minutos = int(round(minutos))
    return f"{minutos // 60}h{minutos % 60:02d}"


# --------------------------------------------------------------------------- #
# Planificación
# --------------------------------------------------------------------------- #
@dataclass(eq=False)  # identidad: dos tareas del mismo punto (instalación/retiro) son distintas
class Tarea:
    punto: Punto
    tipo: str = "muestreo"              # muestreo | instalacion | retiro
    no_antes: datetime | None = None    # retiro: fin de la exposición
    equipos_en_campo: list[str] = field(default_factory=list)  # retiro: códigos a recoger

    @property
    def codigo(self) -> str:
        return self.punto.codigo

    @property
    def duracion(self) -> int:
        return self.punto.duracion_retiro_min if self.tipo == "retiro" else self.punto.duracion_min

    @property
    def prioridad(self) -> int:
        return 0 if self.tipo == "retiro" else self.punto.prioridad

    @property
    def recoge_muestra(self) -> bool:
        return self.tipo in ("muestreo", "retiro")

    def etiqueta(self) -> str:
        p = self.punto
        if self.tipo == "instalacion":
            return f"Instalación de equipo de {p.matriz} ({p.horas_exposicion:g} h)"
        if self.tipo == "retiro":
            eq = f" [{', '.join(self.equipos_en_campo)}]" if self.equipos_en_campo else ""
            return f"Retiro de equipo y muestra de {p.matriz}{eq}"
        return f"Muestreo de {p.matriz}"


def tareas_iniciales(puntos: list[Punto]) -> list[Tarea]:
    return [Tarea(p, "instalacion" if p.horas_exposicion > 0 else "muestreo") for p in puntos]


@dataclass
class Cuadrilla:
    id: str
    vehiculo: Vehiculo
    integrantes: list[Persona]
    portatiles: dict[str, str] = field(default_factory=dict)        # tipo -> código de equipo
    a_instalar: list[tuple[str, str]] = field(default_factory=list)  # (código equipo, código punto)

    @property
    def competencias(self) -> set[str]:
        return {c for p in self.integrantes for c in p.competencias}

    @property
    def conductor(self) -> Persona:
        return next(p for p in self.integrantes if p.conduce)

    @property
    def equipos(self) -> list[str]:
        return sorted(self.portatiles.values()) + [c for c, _ in self.a_instalar]

    def puede_hacer(self, t: Tarea) -> bool:
        p = t.punto
        return (
            p.matriz in self.competencias
            and (self.vehiculo.es_4x4 or not p.requiere_4x4)
            and len(self.integrantes) >= p.personas_min
        )

    def describir(self) -> str:
        nombres = ", ".join(
            f"{x.nombre}{' (cond.)' if x is self.conductor else ''}" for x in self.integrantes
        )
        return f"{self.id} | {self.vehiculo.placa} ({self.vehiculo.descripcion}) | {nombres}"


@dataclass
class Evento:
    actividad: str
    lugar: str
    inicio: datetime
    fin: datetime
    km: float = 0.0
    tarea: Tarea | None = None


def pool_equipos(equipos: list[Equipo], dia: date, instalados: dict[str, str]) -> dict[str, list[Equipo]]:
    pool: dict[str, list[Equipo]] = {}
    for e in sorted(equipos, key=lambda e: e.codigo):
        if e.disponible(dia) and e.codigo not in instalados:
            pool.setdefault(e.tipo, []).append(e)
    return pool


def equipar(t: Tarea, c: Cuadrilla, pool: dict[str, list[Equipo]], aplicar: bool = False) -> bool:
    """¿Hay equipos para esta tarea? Con aplicar=True además los asigna a la cuadrilla.
    Muestreo: un equipo portátil por tipo sirve para todos los puntos del día.
    Instalación: cada punto deja su propio equipo en campo hasta el retiro."""
    p = t.punto
    if t.tipo == "retiro" or not p.equipos:
        return True
    tipos = list(dict.fromkeys(p.equipos))
    if t.tipo == "muestreo":
        tipos = [tp for tp in tipos if tp not in c.portatiles]
    if any(not pool.get(tp) for tp in tipos):
        return False
    if aplicar:
        for tp in tipos:
            cod = pool[tp].pop(0).codigo
            if t.tipo == "muestreo":
                c.portatiles[tp] = cod
            else:
                c.a_instalar.append((cod, p.codigo))
    return True


def formar_cuadrillas(personal: list[Persona], vehiculos: list[Vehiculo], tareas: list[Tarea],
                      cfg: Configuracion, dia: date, carga: dict[str, float],
                      racha: dict[str, tuple[date, int]]) -> list[Cuadrilla]:
    """Arma cuadrillas del día: 1 vehículo + 1 conductor + técnicos que maximicen la
    cobertura de matrices pendientes, prefiriendo a quien lleva menos horas acumuladas.
    Excluye a quien ya cumplió el máximo de días consecutivos."""
    ayer = dia - timedelta(days=1)

    def descansa(p: Persona) -> bool:
        ultimo, n = racha.get(p.nombre, (None, 0))
        return ultimo == ayer and n >= cfg.max_dias_consecutivos

    gente = [p for p in personal if p.disponible(dia) and not descansa(p)]
    autos = [v for v in vehiculos if v.disponible(dia)]
    puntos = [t.punto for t in tareas]
    necesita_4x4 = any(p.requiere_4x4 for p in puntos)
    autos.sort(key=lambda v: (not v.es_4x4 if necesita_4x4 else v.es_4x4, -v.velocidad_kmh))
    matrices_pend = {p.matriz for p in puntos}
    tam_req = max([cfg.tamano_cuadrilla, *(p.personas_min for p in puntos)], default=2)

    cuadrillas: list[Cuadrilla] = []
    for v in autos:
        conductores = [p for p in gente if p.conduce]
        if not conductores:
            break
        # Conductor con menos competencias útiles primero (reserva especialistas), luego menos carga.
        conductores.sort(key=lambda p: (len(set(p.competencias) & matrices_pend), carga.get(p.nombre, 0)))
        chofer = conductores[0]
        gente.remove(chofer)
        equipo = [chofer]
        tam = min(v.capacidad, tam_req)
        while len(equipo) < tam and gente:
            cubiertas = {c for x in equipo for c in x.competencias}
            gente.sort(key=lambda p: (-len((set(p.competencias) & matrices_pend) - cubiertas),
                                      p.conduce, carga.get(p.nombre, 0), -len(p.competencias)))
            equipo.append(gente.pop(0))
        cq = Cuadrilla(f"C{len(cuadrillas) + 1}", v, equipo)
        if cq.competencias & matrices_pend:
            cuadrillas.append(cq)
        else:  # no aporta nada: liberar al personal
            gente.extend(equipo)
    return cuadrillas


def simular(ruta: list[Tarea], cq: Cuadrilla, cfg: Configuracion, dia: date,
            dist: Distancias) -> tuple[list[Evento], datetime, float, str | None]:
    """Línea de tiempo del día. Devuelve (eventos, hora_fin, km, problema);
    problema es None si la ruta cumple jornada, ventanas y preservación."""
    t = datetime.combine(dia, parse_hora(cfg.hora_inicio))
    h_alm = datetime.combine(dia, parse_hora(cfg.hora_almuerzo))
    alm_pendiente = cfg.almuerzo_min > 0
    ev: list[Evento] = []
    km_tot = 0.0
    problema: str | None = None
    muestras: list[tuple[Tarea, datetime]] = []

    def agregar(act: str, lugar: str, minutos: float, km: float = 0.0, tarea: Tarea | None = None) -> None:
        nonlocal t
        fin = t + timedelta(minutes=minutos)
        ev.append(Evento(act, lugar, t, fin, km, tarea))
        t = fin

    def quizas_almorzar(lugar: str) -> None:
        nonlocal alm_pendiente
        if alm_pendiente and t >= h_alm:
            agregar("Almuerzo", lugar, cfg.almuerzo_min)
            alm_pendiente = False

    def esperar_hasta(limite: datetime, motivo: str, lugar: str) -> None:
        nonlocal alm_pendiente
        if t >= limite:
            return
        # Si la espera abarca la hora de almuerzo, se aprovecha para almorzar.
        if alm_pendiente and limite >= h_alm + timedelta(minutes=cfg.almuerzo_min):
            if t < h_alm:
                agregar(f"Espera ({motivo})", lugar, (h_alm - t).total_seconds() / 60)
            agregar("Almuerzo", lugar, cfg.almuerzo_min)
            alm_pendiente = False
        if t < limite:
            agregar(f"Espera ({motivo})", lugar, (limite - t).total_seconds() / 60)

    if cfg.preparacion_min:
        agregar("Preparación de equipos / charla de seguridad", cfg.base_nombre, cfg.preparacion_min)
    actual = BASE
    for tr in ruta:
        p = tr.punto
        quizas_almorzar("En ruta")
        km = dist.km(actual, p.codigo)
        km_tot += km
        agregar(f"Traslado ({km:.1f} km)", f"→ {p.codigo}",
                dist.minutos(actual, p.codigo, cq.vehiculo.velocidad_kmh), km)
        if p.acceso_min:
            agregar("Acceso a pie (ida y vuelta)", p.codigo, p.acceso_min)
        if p.hora_desde:
            esperar_hasta(datetime.combine(dia, parse_hora(p.hora_desde)), "ventana horaria", p.codigo)
        if tr.no_antes:
            esperar_hasta(tr.no_antes, f"fin de exposición {p.horas_exposicion:g} h", p.codigo)
        quizas_almorzar(p.codigo)
        agregar(tr.etiqueta(), f"{p.codigo} - {p.nombre}", tr.duracion, tarea=tr)
        if p.hora_hasta and t > datetime.combine(dia, parse_hora(p.hora_hasta)) and not problema:
            problema = f"{p.codigo}: termina fuera de la ventana {p.hora_desde or '--'}-{p.hora_hasta}"
        if tr.recoge_muestra and p.max_horas_preservacion:
            muestras.append((tr, t))
        actual = p.codigo
    if actual != BASE:
        quizas_almorzar(actual)
    km = dist.km(actual, BASE)
    km_tot += km
    agregar(f"Retorno a base ({km:.1f} km)", cfg.base_nombre,
            dist.minutos(actual, BASE, cq.vehiculo.velocidad_kmh), km)
    for tr, fin_m in muestras:
        if t - fin_m > timedelta(hours=tr.punto.max_horas_preservacion) and not problema:
            problema = f"{tr.codigo}: supera {tr.punto.max_horas_preservacion:g} h de preservación"
    quizas_almorzar(cfg.base_nombre)
    if cfg.descarga_min:
        agregar("Descarga, preservación y registro de muestras", cfg.base_nombre, cfg.descarga_min)
    if t > datetime.combine(dia, parse_hora(cfg.hora_fin)) and not problema:
        problema = "excede la jornada"
    return ev, t, km_tot, problema


def factible(ruta, cq, cfg, dia, dist) -> bool:
    return simular(ruta, cq, cfg, dia, dist)[3] is None


def dos_opt(ruta: list[Tarea], cq, cfg, dia, dist) -> list[Tarea]:
    """Mejora 2-opt: invierte tramos mientras reduzca kilómetros y siga siendo factible."""
    mejor = ruta[:]
    mejor_km = dist.km_ruta([t.codigo for t in mejor])
    mejora = True
    while mejora:
        mejora = False
        for i in range(len(mejor) - 1):
            for j in range(i + 1, len(mejor)):
                nueva = mejor[:i] + mejor[i:j + 1][::-1] + mejor[j + 1:]
                km = dist.km_ruta([t.codigo for t in nueva])
                if km < mejor_km - 1e-6 and factible(nueva, cq, cfg, dia, dist):
                    mejor, mejor_km, mejora = nueva, km, True
    return mejor


@dataclass
class PlanDia:
    dia: date
    cuadrilla: Cuadrilla
    ruta: list[Tarea]
    eventos: list[Evento]
    km: float

    @property
    def minutos(self) -> float:
        return (self.eventos[-1].fin - self.eventos[0].inicio).total_seconds() / 60

    @property
    def litros(self) -> float:
        r = self.cuadrilla.vehiculo.rendimiento_km_l
        return self.km / r if r else 0.0

    def costos(self, cfg: Configuracion) -> tuple[float, float, float]:
        """(combustible, vehículo, viáticos)"""
        return (self.litros * cfg.precio_combustible, self.cuadrilla.vehiculo.costo_dia,
                len(self.cuadrilla.integrantes) * cfg.viatico_persona_dia)


@dataclass
class Resultado:
    plan: list[PlanDia]
    no_asignados: dict[str, str]
    fecha_inicio: date
    fuente_tiempos: str


def planificar(repo: Repositorio, puntos: list[Punto], fecha_inicio: date,
               max_dias: int = 30) -> Resultado:
    cfg = repo.config
    dist = Distancias(puntos, cfg)
    pendientes = sorted(tareas_iniciales(puntos), key=lambda t: (t.prioridad, t.codigo))
    instalados: dict[str, str] = {}              # código equipo -> código punto
    carga: dict[str, float] = {}                 # minutos acumulados por persona
    racha: dict[str, tuple[date, int]] = {}      # último día trabajado, días seguidos
    plan: list[PlanDia] = []
    dia, dias_usados, dias_sin_avance = fecha_inicio, 0, 0

    while pendientes and dias_usados < max_dias:
        if dia.weekday() not in cfg.dias_laborables:
            dia += timedelta(days=1)
            continue
        hoy = [t for t in pendientes if t.no_antes is None or t.no_antes.date() <= dia]
        cuads = formar_cuadrillas(repo.personal, repo.vehiculos, hoy, cfg, dia, carga, racha) if hoy else []
        pool = pool_equipos(repo.equipos, dia, instalados)
        rutas: dict[str, list[Tarea]] = {c.id: [] for c in cuads}
        activas = {c.id for c in cuads}
        # Asignación intercalada (round-robin): cada cuadrilla toma su mejor tarea por turno,
        # lo que reparte el territorio de forma natural entre cuadrillas.
        while activas:
            for c in cuads:
                if c.id not in activas:
                    continue
                ultimo = rutas[c.id][-1].codigo if rutas[c.id] else BASE
                cands = sorted((t for t in hoy if c.puede_hacer(t) and equipar(t, c, pool)),
                               key=lambda t: dist.km(ultimo, t.codigo) * PESO_PRIORIDAD.get(t.prioridad, 1))
                elegido = next((t for t in cands if factible(rutas[c.id] + [t], c, cfg, dia, dist)), None)
                if elegido:
                    equipar(elegido, c, pool, aplicar=True)
                    rutas[c.id].append(elegido)
                    hoy.remove(elegido)
                    pendientes.remove(elegido)
                else:
                    activas.discard(c.id)

        avance = False
        for c in cuads:
            if not rutas[c.id]:
                continue
            avance = True
            ruta = dos_opt(rutas[c.id], c, cfg, dia, dist)
            ev, _, km, _ = simular(ruta, c, cfg, dia, dist)
            pd = PlanDia(dia, c, ruta, ev, km)
            plan.append(pd)
            for cod, pcod in c.a_instalar:
                instalados[cod] = pcod
            for e in ev:
                if e.tarea and e.tarea.tipo == "instalacion":
                    p = e.tarea.punto
                    pendientes.append(Tarea(p, "retiro", e.fin + timedelta(hours=p.horas_exposicion),
                                            [cod for cod, pc in c.a_instalar if pc == p.codigo]))
                elif e.tarea and e.tarea.tipo == "retiro":
                    for cod in e.tarea.equipos_en_campo:
                        instalados.pop(cod, None)  # vuelve a base: disponible desde mañana
            for x in c.integrantes:
                ultimo_dia, n = racha.get(x.nombre, (None, 0))
                racha[x.nombre] = (dia, n + 1 if ultimo_dia == dia - timedelta(days=1) else 1)
                carga[x.nombre] = carga.get(x.nombre, 0) + pd.minutos
        if hoy or avance:
            dias_sin_avance = 0 if avance else dias_sin_avance + 1
        if dias_sin_avance >= 7:  # una semana sin poder asignar nada: son inviables
            break
        dia += timedelta(days=1)
        dias_usados += 1

    no_asignados = {f"{t.codigo} ({t.tipo})": diagnosticar(t, repo, dist, fecha_inicio) for t in pendientes}
    return Resultado(plan, no_asignados, fecha_inicio, dist.fuente)


def diagnosticar(t: Tarea, repo: Repositorio, dist: Distancias, dia: date) -> str:
    cfg, p = repo.config, t.punto
    if t.tipo == "retiro":
        return (f"retiro fuera del horizonte de planificación; equipo(s) "
                f"{', '.join(t.equipos_en_campo) or '-'} quedan instalados")
    if not any(p.matriz in x.competencias and x.activo for x in repo.personal):
        return f"nadie del personal tiene competencia en '{p.matriz}'"
    if p.requiere_4x4 and not any(v.es_4x4 and v.activo for v in repo.vehiculos):
        return "requiere 4x4 y no hay vehículos 4x4 activos"
    for tp in dict.fromkeys(p.equipos):
        unidades = [e for e in repo.equipos if e.tipo == tp and e.activo]
        if not unidades:
            return f"no hay equipos de tipo '{tp}' registrados"
        if not any(vigente(e.calibracion_vence, dia) for e in unidades):
            return f"todos los equipos '{tp}' tienen la calibración vencida"
    if p.hora_desde and p.hora_hasta:
        ventana = (datetime.combine(dia, parse_hora(p.hora_hasta))
                   - datetime.combine(dia, parse_hora(p.hora_desde))).total_seconds() / 60
        if ventana < t.duracion:
            return f"la ventana horaria ({fmt_min(ventana)}) es menor que la duración ({t.duracion} min)"
    vel = max((v.velocidad_kmh for v in repo.vehiculos if v.activo), default=40)
    vuelta = dist.minutos(p.codigo, BASE, vel)
    if p.max_horas_preservacion and vuelta > p.max_horas_preservacion * 60:
        return (f"el retorno a base ({fmt_min(vuelta)}) supera el tiempo de preservación "
                f"({p.max_horas_preservacion:g} h); considere entrega directa al laboratorio")
    jornada = (datetime.combine(dia, parse_hora(cfg.hora_fin))
               - datetime.combine(dia, parse_hora(cfg.hora_inicio))).total_seconds() / 60
    necesario = (cfg.preparacion_min + cfg.descarga_min + cfg.almuerzo_min + p.acceso_min
                 + t.duracion + dist.minutos(BASE, p.codigo, vel) + vuelta)
    if necesario > jornada:
        return (f"no cabe en una jornada ({fmt_min(necesario)} necesarios vs {fmt_min(jornada)}); "
                "considere campamento/pernocte o ampliar horario")
    return ("sin cuadrilla compatible en el horizonte (personal disponible, tamaño de cuadrilla, "
            "vehículos, equipos o ventana horaria)")


# --------------------------------------------------------------------------- #
# Resúmenes y reportes
# --------------------------------------------------------------------------- #
def resumen(res: Resultado, cfg: Configuracion) -> dict:
    r = {"puntos": 0, "retiros": 0, "dias": len({pd.dia for pd in res.plan}), "km": 0.0,
         "litros": 0.0, "combustible": 0.0, "vehiculos": 0.0, "viaticos": 0.0,
         "personas": {}, "flota": {}}
    for pd in res.plan:
        r["puntos"] += sum(t.tipo != "retiro" for t in pd.ruta)
        r["retiros"] += sum(t.tipo == "retiro" for t in pd.ruta)
        comb, veh, via = pd.costos(cfg)
        r["km"] += pd.km
        r["litros"] += pd.litros
        r["combustible"] += comb
        r["vehiculos"] += veh
        r["viaticos"] += via
        for x in pd.cuadrilla.integrantes:
            d = r["personas"].setdefault(x.nombre, {"dias": 0, "min": 0.0})
            d["dias"] += 1
            d["min"] += pd.minutos
        placa = pd.cuadrilla.vehiculo.placa
        d = r["flota"].setdefault(placa, {"dias": 0, "km": 0.0, "litros": 0.0, "costo": 0.0})
        d["dias"] += 1
        d["km"] += pd.km
        d["litros"] += pd.litros
        d["costo"] += comb + veh
    r["total"] = r["combustible"] + r["vehiculos"] + r["viaticos"]
    return r


def imprimir_plan(res: Resultado, cfg: Configuracion) -> None:
    m = cfg.moneda
    if not res.plan:
        print("\nNo se pudo programar ningún punto.")
    dia_actual = None
    for pd in res.plan:
        if pd.dia != dia_actual:
            dia_actual = pd.dia
            print("\n" + "=" * 96)
            print(f" {DIAS_SEMANA[pd.dia.weekday()]} {pd.dia.strftime('%d/%m/%Y')}")
            print("=" * 96)
        c = pd.cuadrilla
        comb, veh, via = pd.costos(cfg)
        print(f"\n  {c.describir()}")
        if c.equipos:
            print(f"  Equipos: {', '.join(c.equipos)}")
        print(f"  Ruta: BASE → {' → '.join(t.codigo for t in pd.ruta)} → BASE   "
              f"| {pd.km:.1f} km | ~{pd.litros:.1f} L | costo {m} {comb + veh + via:,.2f}")
        for e in pd.eventos:
            print(f"    {e.inicio:%H:%M}-{e.fin:%H:%M}  {e.actividad:<52} {e.lugar}")

    r = resumen(res, cfg)
    print("\n" + "-" * 96 + "\n RESUMEN\n" + "-" * 96)
    print(f"  Tiempos: {res.fuente_tiempos}")
    print(f"  Puntos programados: {r['puntos']} (+{r['retiros']} retiros) | Días de campo: {r['dias']}"
          f" | Km: {r['km']:.1f} | Combustible: {r['litros']:.1f} L")
    print(f"  Costos: combustible {m} {r['combustible']:,.2f} + vehículos {m} {r['vehiculos']:,.2f}"
          f" + viáticos {m} {r['viaticos']:,.2f} = TOTAL {m} {r['total']:,.2f}")
    print("  Por persona: " + "; ".join(f"{n} {d['dias']}d/{fmt_min(d['min'])}"
                                         for n, d in sorted(r["personas"].items())))
    print("  Por vehículo: " + "; ".join(f"{p} {d['dias']}d/{d['km']:.0f} km"
                                          for p, d in sorted(r["flota"].items())))
    if res.no_asignados:
        print("\n  NO PROGRAMADOS:")
        for cod, motivo in res.no_asignados.items():
            print(f"   ✗ {cod}: {motivo}")


def exportar_csv(res: Resultado, ruta: Path) -> None:
    with ruta.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Fecha", "Cuadrilla", "Vehículo", "Conductor", "Integrantes", "Equipos",
                    "Inicio", "Fin", "Duración (min)", "Actividad", "Lugar", "Km"])
        for pd in res.plan:
            c = pd.cuadrilla
            for e in pd.eventos:
                w.writerow([pd.dia.isoformat(), c.id, c.vehiculo.placa, c.conductor.nombre,
                            ", ".join(x.nombre for x in c.integrantes), ", ".join(c.equipos),
                            f"{e.inicio:%H:%M}", f"{e.fin:%H:%M}",
                            round((e.fin - e.inicio).total_seconds() / 60),
                            e.actividad, e.lugar, f"{e.km:.1f}".replace(".", ",")])


def exportar_excel(res: Resultado, cfg: Configuracion, ruta: Path) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return False
    wb = Workbook()
    wb.remove(wb.active)
    cab_font = Font(bold=True, color="FFFFFF")
    cab_fill = PatternFill("solid", fgColor="1F6F5C")
    m = cfg.moneda

    def hoja(titulo: str, encabezados: list[str], filas: list[list]) -> None:
        ws = wb.create_sheet(titulo)
        ws.append(encabezados)
        for celda in ws[1]:
            celda.font, celda.fill = cab_font, cab_fill
            celda.alignment = Alignment(vertical="center", wrap_text=True)
        for f in filas:
            ws.append(f)
        ws.freeze_panes = "A2"
        if filas:
            ws.auto_filter.ref = ws.dimensions
        for i, col in enumerate(ws.columns, 1):
            largo = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, largo + 2), 60)
        for fila in ws.iter_rows(min_row=2):
            for c in fila:
                if isinstance(c.value, date):
                    c.number_format = "dd/mm/yyyy"
                elif isinstance(c.value, float):
                    c.number_format = "#,##0.00"

    r = resumen(res, cfg)
    hoja("Resumen", ["Indicador", "Valor"], [
        ["Fecha de inicio", res.fecha_inicio],
        ["Fuente de tiempos", res.fuente_tiempos],
        ["Puntos programados", r["puntos"]],
        ["Retiros de equipos", r["retiros"]],
        ["Días de campo", r["dias"]],
        ["Km totales", round(r["km"], 1)],
        ["Combustible (L)", round(r["litros"], 1)],
        [f"Costo combustible ({m})", round(r["combustible"], 2)],
        [f"Costo vehículos ({m})", round(r["vehiculos"], 2)],
        [f"Viáticos ({m})", round(r["viaticos"], 2)],
        [f"COSTO TOTAL ({m})", round(r["total"], 2)],
        ["Puntos no programados", len(res.no_asignados)],
    ])
    filas_rutas = []
    for pd in res.plan:
        c = pd.cuadrilla
        comb, veh, via = pd.costos(cfg)
        filas_rutas.append([
            pd.dia, DIAS_SEMANA[pd.dia.weekday()], c.id, c.vehiculo.placa, c.conductor.nombre,
            ", ".join(x.nombre for x in c.integrantes), ", ".join(c.equipos),
            " → ".join(f"{t.codigo}{'(R)' if t.tipo == 'retiro' else '(I)' if t.tipo == 'instalacion' else ''}"
                       for t in pd.ruta),
            f"{pd.eventos[0].inicio:%H:%M}", f"{pd.eventos[-1].fin:%H:%M}", round(pd.km, 1),
            round(pd.litros, 1), round(comb, 2), round(veh, 2), round(via, 2), round(comb + veh + via, 2)])
    hoja("Rutas", ["Fecha", "Día", "Cuadrilla", "Vehículo", "Conductor", "Integrantes", "Equipos",
                   "Secuencia (I=instalación, R=retiro)", "Salida", "Fin", "Km", "Litros",
                   f"Combustible {m}", f"Vehículo {m}", f"Viáticos {m}", f"Total {m}"], filas_rutas)
    hoja("Cronograma", ["Fecha", "Cuadrilla", "Vehículo", "Inicio", "Fin", "Min", "Actividad",
                        "Lugar", "Km", "Latitud", "Longitud"],
         [[pd.dia, pd.cuadrilla.id, pd.cuadrilla.vehiculo.placa, f"{e.inicio:%H:%M}", f"{e.fin:%H:%M}",
           round((e.fin - e.inicio).total_seconds() / 60), e.actividad, e.lugar, round(e.km, 1),
           e.tarea.punto.lat if e.tarea else None, e.tarea.punto.lon if e.tarea else None]
          for pd in res.plan for e in pd.eventos])
    hoja("Personal", ["Persona", "Días de campo", "Horas"],
         [[n, d["dias"], round(d["min"] / 60, 1)] for n, d in sorted(r["personas"].items())])
    hoja("Vehículos", ["Placa", "Días", "Km", "Litros", f"Costo {m} (comb.+vehículo)"],
         [[p, d["dias"], round(d["km"], 1), round(d["litros"], 1), round(d["costo"], 2)]
          for p, d in sorted(r["flota"].items())])
    hoja("No programados", ["Punto (tarea)", "Motivo"], [[k, v] for k, v in res.no_asignados.items()])
    wb.save(str(ruta))
    return True


def exportar_hoja_ruta(res: Resultado, repo: Repositorio, ruta: Path) -> None:
    """HTML imprimible: una página por cuadrilla y día, con enlaces de navegación."""
    cfg, e_ = repo.config, html.escape
    base = (cfg.base_lat, cfg.base_lon)
    secciones = []
    for pd in res.plan:
        c = pd.cuadrilla
        comb, veh, via = pd.costos(cfg)
        paradas = [base, *[(t.punto.lat, t.punto.lon) for t in pd.ruta], base]
        ruta_gmaps = "https://www.google.com/maps/dir/" + "/".join(f"{a},{b}" for a, b in paradas)
        filas = []
        for ev in pd.eventos:
            nota, enlace = "", ""
            if ev.tarea:
                p = ev.tarea.punto
                detalles = [f"{p.lat:.5f}, {p.lon:.5f}"]
                if p.hora_desde or p.hora_hasta:
                    detalles.append(f"ventana {p.hora_desde or '--'}–{p.hora_hasta or '--'}")
                if p.max_horas_preservacion and ev.tarea.recoge_muestra:
                    detalles.append(f"preservar ≤ {p.max_horas_preservacion:g} h")
                if p.equipos and ev.tarea.tipo != "retiro":
                    detalles.append("equipos: " + ", ".join(p.equipos))
                if p.requiere_4x4:
                    detalles.append("acceso 4x4")
                nota = f"<div class='nota'>{e_(' · '.join(detalles))}</div>"
                enlace = (f"<a href='https://www.google.com/maps/dir/?api=1&destination={p.lat},{p.lon}"
                          f"&travelmode=driving' target='_blank'>Navegar</a>")
            clase = "muestra" if ev.tarea else ""
            filas.append(f"<tr class='{clase}'><td>{ev.inicio:%H:%M}</td><td>{ev.fin:%H:%M}</td>"
                         f"<td>{e_(ev.actividad)}{nota}</td><td>{e_(ev.lugar)}</td><td>{enlace}</td>"
                         f"<td class='chk'></td></tr>")
        secciones.append(f"""
<section>
  <header>
    <div><h2>{DIAS_SEMANA[pd.dia.weekday()]} {pd.dia:%d/%m/%Y} · Cuadrilla {c.id}</h2>
    <p class="sub">{e_(cfg.base_nombre)} · salida {pd.eventos[0].inicio:%H:%M} · fin {pd.eventos[-1].fin:%H:%M}</p></div>
    <a class="btn" href="{ruta_gmaps}" target="_blank">Ruta completa en Google Maps</a>
  </header>
  <dl>
    <dt>Vehículo</dt><dd>{e_(c.vehiculo.placa)} · {e_(c.vehiculo.descripcion)}</dd>
    <dt>Conductor</dt><dd>{e_(c.conductor.nombre)}</dd>
    <dt>Integrantes</dt><dd>{e_(', '.join(x.nombre for x in c.integrantes))}</dd>
    <dt>Equipos</dt><dd>{e_(', '.join(c.equipos) or '—')}</dd>
    <dt>Recorrido</dt><dd>{pd.km:.1f} km · ~{pd.litros:.1f} L · {cfg.moneda} {comb + veh + via:,.2f}</dd>
  </dl>
  <table>
    <thead><tr><th>Inicio</th><th>Fin</th><th>Actividad</th><th>Lugar</th><th></th><th>✓</th></tr></thead>
    <tbody>{''.join(filas)}</tbody>
  </table>
  <div class="firmas"><div>Responsable de cuadrilla</div><div>Supervisor</div></div>
</section>""")
    doc = f"""<!doctype html>
<html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Hojas de ruta {res.fecha_inicio:%d/%m/%Y}</title>
<style>
  body {{ font: 14px/1.4 system-ui, Segoe UI, sans-serif; color: #1c2b27; background: #f4f6f5; margin: 0; padding: 16px; }}
  section {{ background: #fff; max-width: 960px; margin: 0 auto 24px; padding: 20px 24px; border-radius: 8px;
            box-shadow: 0 1px 3px rgba(0,0,0,.12); }}
  header {{ display: flex; justify-content: space-between; align-items: flex-start; gap: 12px; flex-wrap: wrap; }}
  h1 {{ max-width: 960px; margin: 0 auto 16px; font-size: 20px; }}
  h2 {{ margin: 0; font-size: 18px; color: #1f6f5c; }}
  .sub {{ margin: 2px 0 0; color: #5a6b66; }}
  .btn {{ background: #1f6f5c; color: #fff; padding: 6px 12px; border-radius: 6px; text-decoration: none; font-size: 13px; }}
  dl {{ display: grid; grid-template-columns: 110px 1fr; gap: 2px 12px; margin: 14px 0; }}
  dt {{ font-weight: 600; color: #5a6b66; }} dd {{ margin: 0; }}
  table {{ width: 100%; border-collapse: collapse; }}
  th, td {{ border-bottom: 1px solid #dde3e1; padding: 5px 6px; text-align: left; vertical-align: top; }}
  th {{ background: #eef3f1; font-size: 12px; text-transform: uppercase; color: #5a6b66; }}
  tr.muestra td {{ background: #f0f8f5; font-weight: 600; }}
  .nota {{ font-weight: 400; font-size: 12px; color: #5a6b66; }}
  td.chk {{ width: 28px; border-left: 1px solid #dde3e1; }}
  a {{ color: #1f6f5c; }}
  .firmas {{ display: flex; gap: 40px; margin-top: 48px; }}
  .firmas div {{ flex: 1; border-top: 1px solid #1c2b27; padding-top: 4px; text-align: center; color: #5a6b66; }}
  @media print {{ body {{ background: #fff; padding: 0; }} h1, .btn {{ display: none; }}
                  section {{ box-shadow: none; page-break-after: always; margin: 0; max-width: none; }} }}
  @media (max-width: 600px) {{ section {{ padding: 14px; }} dl {{ grid-template-columns: 1fr; }} }}
</style></head><body>
<h1>Hojas de ruta · programación desde {res.fecha_inicio:%d/%m/%Y}</h1>
{''.join(secciones) or '<p>Sin rutas programadas.</p>'}
</body></html>"""
    ruta.write_text(doc, encoding="utf-8")


def exportar_mapa(res: Resultado, repo: Repositorio, ruta: Path) -> bool:
    try:
        import folium
    except ImportError:
        return False
    cfg = repo.config
    m = folium.Map(location=[cfg.base_lat, cfg.base_lon], zoom_start=10)
    folium.Marker([cfg.base_lat, cfg.base_lon], tooltip=cfg.base_nombre,
                  icon=folium.Icon(color="black", icon="home")).add_to(m)
    colores = ["blue", "red", "green", "purple", "orange", "darkred", "cadetblue", "darkgreen"]
    cache = _leer_cache_osrm() if cfg.usar_osrm else {}
    rectas = 0
    for i, pd in enumerate(res.plan):
        col = colores[i % len(colores)]
        capa = folium.FeatureGroup(name=f"{pd.dia:%d/%m} {pd.cuadrilla.id} ({pd.km:.0f} km)")
        coords = [(cfg.base_lat, cfg.base_lon), *[(t.punto.lat, t.punto.lon) for t in pd.ruta],
                  (cfg.base_lat, cfg.base_lon)]
        linea = trazado_osrm(coords, cfg.osrm_url, cache) if cfg.usar_osrm else None
        if linea:
            folium.PolyLine(linea, color=col, weight=4, opacity=0.8,
                            tooltip=f"{pd.dia:%d/%m} {pd.cuadrilla.id}").add_to(capa)
        else:
            rectas += 1
            folium.PolyLine([list(c) for c in coords], color=col, weight=3, opacity=0.8,
                            dash_array="8", tooltip=f"{pd.dia:%d/%m} {pd.cuadrilla.id} (línea recta)"
                            ).add_to(capa)
        horas = {id(e.tarea): e for e in pd.eventos if e.tarea}
        for n, t in enumerate(pd.ruta, 1):
            p, e = t.punto, horas.get(id(t))
            hora = f" | {e.inicio:%H:%M}-{e.fin:%H:%M}" if e else ""
            folium.CircleMarker([p.lat, p.lon], radius=8, color=col, fill=True, fill_opacity=0.9,
                                tooltip=f"{n}. {p.codigo} - {p.nombre} · {t.etiqueta()}{hora}").add_to(capa)
            folium.Marker([p.lat, p.lon], icon=folium.DivIcon(
                html=f'<div style="font:bold 10px sans-serif;color:#fff;text-align:center;'
                     f'width:16px;margin:-6px 0 0 -8px">{n}</div>')).add_to(capa)
        capa.add_to(m)
    if cfg.usar_osrm:
        _guardar_cache_osrm(cache)
    if rectas:
        print(f"  ⚠ {rectas} ruta(s) sin trazado OSRM en el mapa; se dibujan con línea recta punteada.")
    folium.LayerControl(collapsed=False).add_to(m)
    m.save(str(ruta))
    return True


def exportar_todo(res: Resultado, repo: Repositorio, nombre: str) -> None:
    CARPETA_SALIDAS.mkdir(exist_ok=True)
    base = CARPETA_SALIDAS / nombre
    exportar_csv(res, base.with_suffix(".csv"))
    print(f"  CSV:           {base.with_suffix('.csv')}")
    if exportar_excel(res, repo.config, base.with_suffix(".xlsx")):
        print(f"  Excel:         {base.with_suffix('.xlsx')}")
    else:
        print("  (Instale 'openpyxl' para exportar a Excel)")
    hoja = base.with_name(base.name + "_hojas_de_ruta.html")
    exportar_hoja_ruta(res, repo, hoja)
    print(f"  Hojas de ruta: {hoja}")
    mapa = base.with_name(base.name + "_mapa.html")
    if exportar_mapa(res, repo, mapa):
        print(f"  Mapa:          {mapa}")
    else:
        print("  (Instale 'folium' para generar el mapa)")


def guardar_historial(res: Resultado, repo: Repositorio, registrar: bool) -> Path:
    """Guarda la programación en el historial y, opcionalmente, registra la fecha
    programada como último monitoreo de cada punto (para el cálculo de frecuencias)."""
    CARPETA_HISTORIAL.mkdir(exist_ok=True)
    r = resumen(res, repo.config)
    datos = {
        "generado": datetime.now().isoformat(timespec="seconds"),
        "inicio": res.fecha_inicio.isoformat(),
        "fuente_tiempos": res.fuente_tiempos,
        "resumen": {k: r[k] for k in ("puntos", "retiros", "dias", "km", "litros", "total")},
        "rutas": [{
            "fecha": pd.dia.isoformat(), "cuadrilla": pd.cuadrilla.id,
            "vehiculo": pd.cuadrilla.vehiculo.placa,
            "integrantes": [x.nombre for x in pd.cuadrilla.integrantes],
            "equipos": pd.cuadrilla.equipos, "km": round(pd.km, 1),
            "tareas": [f"{t.codigo}:{t.tipo}" for t in pd.ruta],
            "eventos": [[f"{e.inicio:%H:%M}", f"{e.fin:%H:%M}", e.actividad, e.lugar] for e in pd.eventos],
        } for pd in res.plan],
        "no_programados": res.no_asignados,
    }
    ruta = CARPETA_HISTORIAL / f"programacion_{res.fecha_inicio:%Y%m%d}_{datetime.now():%H%M%S}.json"
    ruta.write_text(json.dumps(datos, indent=1, ensure_ascii=False), encoding="utf-8")
    if registrar:
        for pd in res.plan:
            for t in pd.ruta:
                if t.tipo != "retiro":
                    t.punto.ultimo_monitoreo = pd.dia.isoformat()
        repo.guardar()
    return ruta


def matriz_tiempos(repo: Repositorio) -> None:
    pts = [p for p in repo.puntos if p.activo]
    if not pts:
        print("No hay puntos registrados.")
        return
    dist = Distancias(pts, repo.config)
    vels = [v.velocidad_kmh for v in repo.vehiculos if v.activo] or [40.0]
    vel = sum(vels) / len(vels)
    codigos = [BASE, *(p.codigo for p in pts)]
    etiqueta = lambda c: "BASE" if c == BASE else c
    ancho = max(8, *(len(etiqueta(c)) + 1 for c in codigos))
    for titulo, valor in (("Tiempo de traslado (min), fila = origen, columna = destino",
                           lambda a, b: dist.minutos(a, b, vel)),
                          ("Distancia (km)", dist.km)):
        print(f"\n{titulo}:")
        print(" " * ancho + "".join(f"{etiqueta(c):>{ancho}}" for c in codigos))
        for a in codigos:
            print(f"{etiqueta(a):<{ancho}}" + "".join(
                f"{valor(a, b):>{ancho}.0f}" for b in codigos))
    if not dist.rutas:
        print(f"\n(Tramos estimados a {vel:.0f} km/h, velocidad promedio de la flota.)")


def alertas(repo: Repositorio) -> None:
    hoy = date.today()
    limite = hoy + timedelta(days=DIAS_ALERTA)

    def estado(hasta: str) -> str | None:
        if not hasta:
            return None
        f = date.fromisoformat(hasta)
        if f < hoy:
            return f"VENCIDO el {f:%d/%m/%Y}"
        if f <= limite:
            return f"vence en {(f - hoy).days} días ({f:%d/%m/%Y})"
        return None

    print(f"\n--- ALERTAS (próximos {DIAS_ALERTA} días) ---")
    avisos = [(f"Personal  {p.nombre}", "habilitación (EMO/SCTR/inducción)", estado(p.habilitado_hasta))
              for p in repo.personal if p.activo]
    avisos += [(f"Vehículo  {v.placa}", "documentos (SOAT/revisión técnica)", estado(v.documentos_hasta))
               for v in repo.vehiculos if v.activo]
    avisos += [(f"Equipo    {e.codigo}", f"calibración ({e.tipo})", estado(e.calibracion_vence))
               for e in repo.equipos if e.activo]
    avisos = [a for a in avisos if a[2]]
    for quien, que, est in avisos:
        print(f"  {'✗' if 'VENCIDO' in est else '!'} {quien:<28} {que:<36} {est}")
    if not avisos:
        print("  Sin vencimientos próximos de habilitaciones, documentos ni calibraciones.")

    print("\n--- PUNTOS SEGÚN FRECUENCIA ---")
    filas = []
    for p in repo.puntos:
        if not p.activo:
            continue
        f = p.proxima_fecha()
        if f is None:
            filas.append((date.max, p, "realizado (monitoreo único)"))
        elif f == date.min:
            filas.append((f, p, "PENDIENTE (sin registro previo)"))
        else:
            d = (f - hoy).days
            filas.append((f, p, f"VENCIDO hace {-d} días" if d < 0 else "vence HOY" if d == 0 else
                          f"vence en {d} días ({f:%d/%m/%Y})"))
    for _, p, est in sorted(filas, key=lambda x: x[0]):
        print(f"  {p.codigo:<8} {p.frecuencia:<11} último: {p.ultimo_monitoreo or '—':<10}  {est}")


# --------------------------------------------------------------------------- #
# Entrada por consola
# --------------------------------------------------------------------------- #
def pedir(txt: str, tipo=str, defecto=None):
    sufijo = f" [{defecto}]" if defecto not in (None, "") else ""
    while True:
        r = input(f"{txt}{sufijo}: ").strip()
        if not r and defecto is not None:
            return defecto
        try:
            if tipo is bool:
                if r.lower() in SI:
                    return True
                if r.lower() in ("n", "no", "0", "false"):
                    return False
                raise ValueError
            return tipo(r.replace(",", ".") if tipo is float else r)
        except ValueError:
            print("  Valor no válido, intente de nuevo.")


def pedir_hora(txt: str, defecto: str = "", opcional: bool = False) -> str:
    extra = " ('-' = sin límite)" if opcional else ""
    while True:
        r = input(f"{txt}{extra}{f' [{defecto}]' if defecto else ''}: ").strip()
        if not r:
            return defecto
        if opcional and r == "-":
            return ""
        try:
            return f"{parse_hora(r):%H:%M}"
        except ValueError:
            print("  Use el formato HH:MM.")


def pedir_fecha(txt: str, defecto: str = "") -> str:
    while True:
        r = input(f"{txt} (AAAA-MM-DD, '-' = ninguna){f' [{defecto}]' if defecto else ''}: ").strip()
        if not r:
            return defecto
        if r == "-":
            return ""
        try:
            return date.fromisoformat(r).isoformat()
        except ValueError:
            print("  Fecha inválida.")


def pedir_lista(txt: str, opciones: list[str] | None = None, defecto: list[str] | None = None,
                sugerencias: list[str] | None = None) -> list[str]:
    if opciones or sugerencias:
        print(f"  {'Opciones' if opciones else 'Registrados'}: {', '.join(opciones or sugerencias)}")
    sufijo = " [" + ",".join(defecto) + "]" if defecto else ""
    r = input(f"{txt} (separado por comas, '-' = ninguno){sufijo}: ").strip()
    if not r:
        return defecto or []
    if r == "-":
        return []
    items = [x.strip().lower() for x in r.split(",") if x.strip()]
    if opciones:
        malos = [x for x in items if x not in opciones]
        if malos:
            print(f"  Ignorados (no válidos): {', '.join(malos)}")
        items = [x for x in items if x in opciones]
    return items


def pedir_fechas(txt: str, defecto: list[str]) -> list[str]:
    r = input(f"{txt} (AAAA-MM-DD, separadas por comas; rango con '..'; '-' = ninguna)"
              f"{' [' + ','.join(defecto) + ']' if defecto else ''}: ").strip()
    if not r:
        return defecto
    if r == "-":
        return []
    fechas: list[str] = []
    for parte in r.split(","):
        parte = parte.strip()
        try:
            if ".." in parte:
                a, b = (date.fromisoformat(x.strip()) for x in parte.split(".."))
                while a <= b:
                    fechas.append(a.isoformat())
                    a += timedelta(days=1)
            elif parte:
                fechas.append(date.fromisoformat(parte).isoformat())
        except ValueError:
            print(f"  Fecha inválida ignorada: {parte}")
    return sorted(set(fechas))


def elegir(items: list, etiqueta) -> int | None:
    if not items:
        print("  (lista vacía)")
        return None
    for i, x in enumerate(items, 1):
        print(f"  {i:>3}. {etiqueta(x)}")
    n = pedir("Número (0 cancela)", int, 0)
    return n - 1 if 1 <= n <= len(items) else None


# --------------------------------------------------------------------------- #
# Importación / exportación CSV genérica
# --------------------------------------------------------------------------- #
def _a_texto(v) -> str:
    if isinstance(v, bool):
        return "si" if v else "no"
    if isinstance(v, list):
        return "|".join(map(str, v))
    if isinstance(v, float):
        return f"{v:g}".replace(".", ",")
    return str(v)


def _convertir(valor: str, tipo: str):
    v = valor.strip()
    if tipo == "bool":
        return v.lower() in SI
    if tipo == "int":
        return int(float(v.replace(",", ".")))
    if tipo == "float":
        return float(v.replace(",", "."))
    if tipo.startswith("list"):
        return [x.strip().lower() for x in v.replace(";", "|").split("|") if x.strip()]
    return v


def exportar_entidad_csv(lista: list, cls, nombre: str) -> None:
    CARPETA_SALIDAS.mkdir(exist_ok=True)
    ruta = CARPETA_SALIDAS / f"{nombre}.csv"
    nombres = [f.name for f in fields(cls)]
    with ruta.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(nombres)
        for x in lista:
            w.writerow([_a_texto(getattr(x, n)) for n in nombres])
    print(f"  Exportado ({len(lista)} registros): {ruta}")
    print("  Edítelo en Excel y use 'Importar CSV' para cargar cambios (se actualiza por clave).")


def importar_csv(repo: Repositorio, lista: list, cls, clave: str) -> None:
    ruta = Path(pedir("Ruta del archivo CSV", str).strip('"'))
    if not ruta.exists():
        print("  Archivo no encontrado.")
        return
    texto = ruta.read_text(encoding="utf-8-sig")
    if not texto.strip():
        print("  Archivo vacío.")
        return
    dialecto = csv.Sniffer().sniff(texto.splitlines()[0], delimiters=",;\t")
    campos = {f.name: f for f in fields(cls)}
    obligatorios = [f.name for f in fields(cls) if f.default is MISSING and f.default_factory is MISSING]
    nuevos = actualizados = errores = 0
    for n, fila in enumerate(csv.DictReader(texto.splitlines(), dialect=dialecto), 2):
        fila = {k.strip().lower(): (v or "") for k, v in fila.items() if k}
        try:
            faltan = [c for c in obligatorios if not fila.get(c, "").strip()]
            if faltan:
                raise ValueError(f"faltan columnas/valores: {', '.join(faltan)}")
            obj = cls(**{k: _convertir(v, campos[k].type) for k, v in fila.items()
                         if k in campos and v.strip()})
            if isinstance(obj, Punto):
                obj.codigo, obj.matriz = obj.codigo.upper(), obj.matriz.lower()
                if obj.matriz not in MATRICES:
                    raise ValueError(f"matriz '{obj.matriz}' no válida")
                if obj.frecuencia not in FRECUENCIAS:
                    raise ValueError(f"frecuencia '{obj.frecuencia}' no válida")
            elif isinstance(obj, Vehiculo):
                obj.placa = obj.placa.upper()
            elif isinstance(obj, Equipo):
                obj.codigo, obj.tipo = obj.codigo.upper(), obj.tipo.lower()
        except (ValueError, TypeError) as e:
            print(f"  Fila {n}: {e}; omitida.")
            errores += 1
            continue
        idx = next((i for i, x in enumerate(lista) if getattr(x, clave) == getattr(obj, clave)), None)
        if idx is None:
            lista.append(obj)
            nuevos += 1
        else:
            lista[idx] = obj
            actualizados += 1
    repo.guardar()
    print(f"  Nuevos: {nuevos} | Actualizados: {actualizados} | Con error: {errores}")


# --------------------------------------------------------------------------- #
# Menús
# --------------------------------------------------------------------------- #
def fmt_persona(p: Persona) -> str:
    estado = "" if p.activo else " [INACTIVO]"
    nd = f" | no disp.: {len(p.no_disponible)} día(s)" if p.no_disponible else ""
    hab = f" | habilitado hasta {p.habilitado_hasta}" if p.habilitado_hasta else ""
    return (f"{p.nombre:<22} {p.cargo:<14} {'🚗 conduce' if p.conduce else '         '}  "
            f"[{', '.join(p.competencias)}]{hab}{nd}{estado}")


def fmt_vehiculo(v: Vehiculo) -> str:
    estado = "" if v.activo else " [INACTIVO]"
    nd = f" | no disp.: {len(v.no_disponible)} día(s)" if v.no_disponible else ""
    doc = f" | docs hasta {v.documentos_hasta}" if v.documentos_hasta else ""
    return (f"{v.placa:<10} {v.descripcion:<18} {v.capacidad} plazas  "
            f"{'4x4' if v.es_4x4 else '4x2'}  {v.velocidad_kmh:.0f} km/h  "
            f"{v.rendimiento_km_l:.1f} km/L  {v.costo_dia:g}/día{doc}{nd}{estado}")


def fmt_equipo(e: Equipo) -> str:
    estado = "" if e.activo else " [INACTIVO]"
    cal = f"calibración hasta {e.calibracion_vence}" if e.calibracion_vence else "sin fecha de calibración"
    return f"{e.codigo:<12} {e.tipo:<18} {e.descripcion[:26]:<26} {cal}{estado}"


def fmt_punto(p: Punto) -> str:
    estado = "" if p.activo else " [INACTIVO]"
    extras = []
    if p.requiere_4x4:
        extras.append("4x4")
    if p.acceso_min:
        extras.append(f"+{p.acceso_min} min a pie")
    if p.hora_desde or p.hora_hasta:
        extras.append(f"{p.hora_desde or '--'}-{p.hora_hasta or '--'}")
    if p.max_horas_preservacion:
        extras.append(f"preserv. {p.max_horas_preservacion:g} h")
    if p.horas_exposicion:
        extras.append(f"exposición {p.horas_exposicion:g} h")
    if p.equipos:
        extras.append("eq: " + ",".join(p.equipos))
    return (f"{p.codigo:<8} {p.nombre[:28]:<28} {p.matriz:<10} ({p.lat:.5f}, {p.lon:.5f})  "
            f"{p.duracion_min} min  P{p.prioridad}  {p.frecuencia}  {'  '.join(extras)}{estado}")


def form_persona(p: Persona | None = None) -> Persona:
    p = p or Persona(nombre="")
    return Persona(
        nombre=pedir("Nombre", str, p.nombre or None),
        cargo=pedir("Cargo", str, p.cargo),
        competencias=pedir_lista("Competencias (matrices)", MATRICES, p.competencias),
        conduce=pedir("¿Conduce? (s/n)", bool, p.conduce),
        activo=pedir("¿Activo? (s/n)", bool, p.activo),
        habilitado_hasta=pedir_fecha("Habilitación (EMO/SCTR/inducción) vigente hasta", p.habilitado_hasta),
        no_disponible=pedir_fechas("Días NO disponible", p.no_disponible),
    )


def form_vehiculo(v: Vehiculo | None = None) -> Vehiculo:
    v = v or Vehiculo(placa="")
    return Vehiculo(
        placa=pedir("Placa", str, v.placa or None).upper(),
        descripcion=pedir("Descripción", str, v.descripcion),
        capacidad=pedir("Capacidad (personas)", int, v.capacidad),
        es_4x4=pedir("¿Es 4x4? (s/n)", bool, v.es_4x4),
        velocidad_kmh=pedir("Velocidad promedio (km/h, si no hay OSRM)", float, v.velocidad_kmh),
        rendimiento_km_l=pedir("Rendimiento (km/L)", float, v.rendimiento_km_l),
        costo_dia=pedir("Costo por día de uso (alquiler/depreciación)", float, v.costo_dia),
        activo=pedir("¿Activo? (s/n)", bool, v.activo),
        documentos_hasta=pedir_fecha("SOAT / revisión técnica vigente hasta", v.documentos_hasta),
        no_disponible=pedir_fechas("Días NO disponible (mantenimiento)", v.no_disponible),
    )


def form_equipo(e: Equipo | None = None, tipos: list[str] | None = None) -> Equipo:
    e = e or Equipo(codigo="", tipo="")
    if tipos:
        print(f"  Tipos registrados: {', '.join(tipos)}")
    return Equipo(
        codigo=pedir("Código / serie", str, e.codigo or None).upper(),
        tipo=pedir("Tipo (p. ej. multiparametro, sonometro, muestreador_pm10)", str, e.tipo or None).lower(),
        descripcion=pedir("Descripción / modelo", str, e.descripcion),
        calibracion_vence=pedir_fecha("Calibración vigente hasta", e.calibracion_vence),
        activo=pedir("¿Activo? (s/n)", bool, e.activo),
        no_disponible=pedir_fechas("Días NO disponible (mantenimiento/préstamo)", e.no_disponible),
    )


def form_punto(p: Punto | None = None, tipos_equipo: list[str] | None = None) -> Punto:
    nuevo = p is None
    p = p or Punto(codigo="", nombre="", lat=0.0, lon=0.0)
    print(f"  Matrices: {', '.join(MATRICES)}")
    matriz = pedir("Matriz", str, p.matriz).lower()
    while matriz not in MATRICES:
        matriz = pedir("Matriz no válida, repita", str).lower()
    q = Punto(
        codigo=pedir("Código", str, p.codigo or None).upper(),
        nombre=pedir("Nombre / descripción", str, p.nombre or None),
        lat=pedir("Latitud (decimal, ej. -12.0464)", float, None if nuevo else p.lat),
        lon=pedir("Longitud (decimal, ej. -77.0428)", float, None if nuevo else p.lon),
        matriz=matriz,
        duracion_min=pedir("Duración del muestreo o instalación (min)", int, p.duracion_min),
        requiere_4x4=pedir("¿Requiere 4x4? (s/n)", bool, p.requiere_4x4),
        acceso_min=pedir("Caminata ida+vuelta (min)", int, p.acceso_min),
        prioridad=pedir("Prioridad (1 alta, 2 media, 3 baja)", int, p.prioridad),
        personas_min=pedir("Personas mínimas", int, p.personas_min),
        activo=pedir("¿Activo? (s/n)", bool, p.activo),
    )
    # Conservar opciones avanzadas y permitir editarlas.
    for f in ("hora_desde", "hora_hasta", "max_horas_preservacion", "equipos", "horas_exposicion",
              "duracion_retiro_min", "frecuencia", "ultimo_monitoreo"):
        setattr(q, f, getattr(p, f))
    if pedir("¿Editar opciones avanzadas (ventana, preservación, equipos, exposición, frecuencia)? (s/n)",
             bool, False):
        q.hora_desde = pedir_hora("Muestrear desde (HH:MM)", q.hora_desde, opcional=True)
        q.hora_hasta = pedir_hora("Muestrear hasta (HH:MM)", q.hora_hasta, opcional=True)
        q.max_horas_preservacion = pedir("Horas máx. de la toma hasta llegar a base (0 = sin límite)",
                                         float, q.max_horas_preservacion)
        q.equipos = pedir_lista("Tipos de equipo necesarios", None, q.equipos, tipos_equipo)
        q.horas_exposicion = pedir("Horas de exposición con equipo instalado (0 = muestreo puntual)",
                                   float, q.horas_exposicion)
        if q.horas_exposicion:
            q.duracion_retiro_min = pedir("Duración del retiro (min)", int, q.duracion_retiro_min)
        print(f"  Frecuencias: {', '.join(FRECUENCIAS)}")
        fr = pedir("Frecuencia", str, q.frecuencia).lower()
        q.frecuencia = fr if fr in FRECUENCIAS else q.frecuencia
        q.ultimo_monitoreo = pedir_fecha("Fecha del último monitoreo", q.ultimo_monitoreo)
    return q


def menu_crud(titulo: str, lista: list, fmt, form, clave: str, repo: Repositorio, cls, nombre_csv: str) -> None:
    while True:
        print(f"\n--- {titulo} ({len(lista)}) ---")
        for x in lista:
            print("  " + fmt(x))
        print("\n 1) Agregar  2) Editar  3) Eliminar  4) Importar CSV  5) Exportar CSV  0) Volver")
        op = input("> ").strip()
        if op == "1":
            nuevo = form()
            if any(getattr(x, clave) == getattr(nuevo, clave) for x in lista):
                print("  Ya existe un registro con esa clave.")
            else:
                lista.append(nuevo)
                repo.guardar()
        elif op == "2":
            i = elegir(lista, fmt)
            if i is not None:
                lista[i] = form(lista[i])
                repo.guardar()
        elif op == "3":
            i = elegir(lista, fmt)
            if i is not None and pedir(f"¿Eliminar {getattr(lista[i], clave)}? (s/n)", bool, False):
                lista.pop(i)
                repo.guardar()
        elif op == "4":
            importar_csv(repo, lista, cls, clave)
        elif op == "5":
            exportar_entidad_csv(lista, cls, nombre_csv)
        elif op == "0":
            return


def menu_config(repo: Repositorio) -> None:
    c = repo.config
    print("\n--- Configuración (Enter mantiene el valor actual) ---")
    c.base_nombre = pedir("Nombre de la base", str, c.base_nombre)
    c.base_lat = pedir("Latitud base", float, c.base_lat)
    c.base_lon = pedir("Longitud base", float, c.base_lon)
    c.hora_inicio = pedir_hora("Hora de salida (HH:MM)", c.hora_inicio)
    c.hora_fin = pedir_hora("Hora límite de fin de jornada (HH:MM)", c.hora_fin)
    c.preparacion_min = pedir("Preparación en base al inicio (min)", int, c.preparacion_min)
    c.descarga_min = pedir("Descarga/registro al volver (min)", int, c.descarga_min)
    c.hora_almuerzo = pedir_hora("Hora de almuerzo (HH:MM)", c.hora_almuerzo)
    c.almuerzo_min = pedir("Duración almuerzo (min, 0 = sin almuerzo)", int, c.almuerzo_min)
    c.usar_osrm = pedir("¿Usar OSRM para rutas reales por carretera? (s/n)", bool, c.usar_osrm)
    if c.usar_osrm:
        c.osrm_url = pedir("URL del servidor OSRM", str, c.osrm_url)
        c.factor_tiempo_osrm = pedir("Factor sobre tiempos OSRM (1.0 = tal cual; 1.2–1.5 tráfico/trocha)",
                                     float, c.factor_tiempo_osrm)
    c.factor_ruta = pedir("Factor de ruta si no hay OSRM (1.2 llano – 1.6 sierra/trocha)",
                          float, c.factor_ruta)
    c.tamano_cuadrilla = pedir("Personas por cuadrilla", int, c.tamano_cuadrilla)
    c.max_dias_consecutivos = pedir("Máximo de días de campo consecutivos por persona", int,
                                    c.max_dias_consecutivos)
    print("  Días laborables: 0=Lun 1=Mar 2=Mié 3=Jue 4=Vie 5=Sáb 6=Dom")
    dl = pedir("Días laborables", str, ",".join(map(str, c.dias_laborables)))
    c.dias_laborables = sorted({int(x) for x in dl.split(",") if x.strip().isdigit() and int(x) < 7})
    c.moneda = pedir("Moneda", str, c.moneda)
    c.precio_combustible = pedir("Precio del combustible por litro", float, c.precio_combustible)
    c.viatico_persona_dia = pedir("Viático por persona y día", float, c.viatico_persona_dia)
    repo.guardar()
    print("  Configuración guardada.")


def seleccionar_puntos(repo: Repositorio) -> list[Punto]:
    activos = [p for p in repo.puntos if p.activo]
    print(f"\n  Puntos activos: {len(activos)}")
    print("  1) Todos los activos\n  2) Solo los que vencen hasta una fecha (según frecuencia)"
          "\n  3) Por matriz\n  4) Códigos específicos")
    op = pedir("Opción", str, "1")
    if op == "2":
        hasta = pedir("Incluir puntos que vencen hasta (AAAA-MM-DD)", date.fromisoformat,
                      date.today() + timedelta(days=30))
        activos = [p for p in activos if (f := p.proxima_fecha()) is not None and f <= hasta]
    elif op == "3":
        filtro = pedir_lista("Matrices", MATRICES)
        activos = [p for p in activos if p.matriz in filtro]
    elif op == "4":
        sel = {x.strip().upper() for x in input("Códigos (separados por comas): ").split(",")}
        activos = [p for p in activos if p.codigo in sel]
    print(f"  Seleccionados: {len(activos)} punto(s)")
    return activos


def menu_planificar(repo: Repositorio) -> None:
    puntos = seleccionar_puntos(repo)
    if not puntos:
        print("  No hay puntos que programar.")
        return
    inicio = pedir("Fecha de inicio (AAAA-MM-DD)", date.fromisoformat, date.today() + timedelta(days=1))
    max_dias = pedir("Máximo de días de campo", int, 30)
    res = planificar(repo, puntos, inicio, max_dias)
    imprimir_plan(res, repo.config)
    if not res.plan:
        return
    if pedir("\n¿Exportar (CSV, Excel, hojas de ruta, mapa)? (s/n)", bool, True):
        exportar_todo(res, repo, f"programacion_{inicio:%Y%m%d}")
    if pedir("¿Guardar en el historial? (s/n)", bool, True):
        registrar = pedir("¿Registrar las fechas programadas como último monitoreo de cada punto? (s/n)",
                          bool, False)
        print(f"  Historial: {guardar_historial(res, repo, registrar)}")


def menu_historial() -> None:
    archivos = sorted(CARPETA_HISTORIAL.glob("programacion_*.json"), reverse=True) \
        if CARPETA_HISTORIAL.exists() else []
    datos = []
    for a in archivos:
        try:
            datos.append((a, json.loads(a.read_text(encoding="utf-8"))))
        except (OSError, ValueError):
            pass
    fmt = lambda x: (f"{x[1]['generado'][:16].replace('T', ' ')}  inicio {x[1]['inicio']}  "
                     f"{x[1]['resumen']['puntos']} puntos  {x[1]['resumen']['dias']} días  "
                     f"{x[1]['resumen']['km']:.0f} km  total {x[1]['resumen']['total']:,.2f}")
    print("\n--- HISTORIAL DE PROGRAMACIONES ---")
    i = elegir(datos, fmt)
    if i is None:
        return
    d = datos[i][1]
    for r in d["rutas"]:
        print(f"\n  {r['fecha']}  {r['cuadrilla']} | {r['vehiculo']} | {', '.join(r['integrantes'])}"
              f" | {r['km']} km")
        if r["equipos"]:
            print(f"  Equipos: {', '.join(r['equipos'])}")
        for ini, fin, act, lugar in r["eventos"]:
            print(f"    {ini}-{fin}  {act:<52} {lugar}")
    for k, v in d.get("no_programados", {}).items():
        print(f"   ✗ {k}: {v}")


def cargar_demo(repo: Repositorio) -> None:
    hoy = date.today()
    en = lambda dias: (hoy + timedelta(days=dias)).isoformat()
    repo.config = Configuracion(base_nombre="Base Lima", base_lat=-12.0464, base_lon=-77.0428)
    repo.personal = [
        Persona("Ana Quispe", "Especialista", ["agua", "sedimento", "biologico"], habilitado_hasta=en(180)),
        Persona("Luis Rojas", "Técnico", ["agua", "suelo"], conduce=True, habilitado_hasta=en(12)),
        Persona("Carla Mendoza", "Especialista", ["aire", "ruido"], habilitado_hasta=en(200)),
        Persona("Jorge Salas", "Conductor", ["ruido"], conduce=True, habilitado_hasta=en(300)),
        Persona("María Torres", "Técnico", ["agua", "efluente", "suelo"], habilitado_hasta=en(150)),
        Persona("Pedro Huamán", "Técnico", ["aire", "suelo"], conduce=True, habilitado_hasta=en(90)),
        Persona("Rosa Flores", "Técnico", ["agua", "aire"], no_disponible=[en(2)], habilitado_hasta=en(60)),
    ]
    repo.vehiculos = [
        Vehiculo("ABC-123", "Hilux 4x4", 5, True, 45, 9.5, 250, documentos_hasta=en(120)),
        Vehiculo("XYZ-789", "Camioneta 4x2", 4, False, 55, 11.0, 180, documentos_hasta=en(20)),
    ]
    repo.equipos = [
        Equipo("MULTI-01", "multiparametro", "YSI ProDSS", en(90)),
        Equipo("MULTI-02", "multiparametro", "Hanna HI98194", en(40)),
        Equipo("SON-01", "sonometro", "Clase 1", en(200)),
        Equipo("SON-02", "sonometro", "Clase 1", en(-5)),
        Equipo("PM10-01", "muestreador_pm10", "Hi-Vol PM10", en(100)),
        Equipo("PM10-02", "muestreador_pm10", "Hi-Vol PM10", en(25)),
        Equipo("BARR-01", "barreno", "Barreno manual acero inox.", ""),
    ]
    agua = ["multiparametro"]
    repo.puntos = [
        Punto("AG-01", "Río Rímac - Chosica", -11.9386, -76.6970, "agua", 60, False, 10, 1, equipos=agua,
              max_horas_preservacion=6, frecuencia="mensual", ultimo_monitoreo=en(-35)),
        Punto("AG-02", "Río Rímac - Ricardo Palma", -11.9180, -76.6600, "agua", 60, False, 5, 1,
              equipos=agua, frecuencia="mensual", ultimo_monitoreo=en(-28)),
        Punto("AG-03", "Quebrada Santa Eulalia", -11.8950, -76.6650, "agua", 50, True, 25, 2,
              equipos=agua, frecuencia="trimestral"),
        Punto("AG-04", "Río Chillón - Trapiche", -11.8130, -76.9900, "agua", 60, True, 15, 2,
              equipos=agua, frecuencia="mensual", ultimo_monitoreo=en(-40)),
        Punto("SD-01", "Sedimento Rímac - Huachipa", -12.0170, -76.9300, "sedimento", 45, False, 10, 2,
              frecuencia="semestral"),
        Punto("EF-01", "Efluente planta industrial", -12.0600, -77.1200, "efluente", 40, False, 0, 1,
              equipos=agua, hora_desde="09:00", hora_hasta="12:00", frecuencia="mensual",
              ultimo_monitoreo=en(-31)),
        Punto("AI-01", "Calidad de aire PM10 - Ate", -12.0260, -76.9180, "aire", 45, False, 0, 1,
              equipos=["muestreador_pm10"], horas_exposicion=24, duracion_retiro_min=30,
              frecuencia="trimestral"),
        Punto("AI-02", "Calidad de aire PM10 - Callao", -12.0500, -77.1300, "aire", 45, False, 0, 2,
              equipos=["muestreador_pm10"], horas_exposicion=24, duracion_retiro_min=30,
              frecuencia="trimestral"),
        Punto("RU-01", "Ruido ambiental - Cercado", -12.0550, -77.0350, "ruido", 30, False, 0, 3,
              equipos=["sonometro"], hora_desde="07:00", hora_hasta="22:00", frecuencia="semestral"),
        Punto("RU-02", "Ruido ambiental - Villa El Salvador", -12.2130, -76.9370, "ruido", 30, False, 0, 3,
              equipos=["sonometro"], frecuencia="semestral"),
        Punto("SU-01", "Suelo - Lurín", -12.2750, -76.8700, "suelo", 60, True, 20, 2,
              equipos=["barreno"], frecuencia="anual"),
        Punto("SU-02", "Suelo - Pachacámac", -12.2300, -76.8600, "suelo", 60, False, 10, 2,
              equipos=["barreno"], frecuencia="anual"),
        Punto("AG-05", "Río Lurín - Cieneguilla", -12.1100, -76.8100, "agua", 60, False, 10, 2,
              equipos=agua, frecuencia="mensual", ultimo_monitoreo=en(-10)),
        Punto("BI-01", "Hidrobiología - Lomas de Lúcumo", -12.2400, -76.8300, "biologico", 90, True, 30, 3,
              frecuencia="anual", ultimo_monitoreo=en(-380)),
    ]
    repo.guardar()
    print("  Datos de ejemplo cargados.")


def main() -> None:
    for flujo in (sys.stdout, sys.stderr):  # evita errores con tildes/símbolos en consolas no UTF-8
        try:
            flujo.reconfigure(errors="replace")
        except AttributeError:
            pass
    repo = Repositorio()
    if "--demo" in sys.argv:
        cargar_demo(repo)
        alertas(repo)
        res = planificar(repo, [p for p in repo.puntos if p.activo], date.today() + timedelta(days=1))
        imprimir_plan(res, repo.config)
        print()
        exportar_todo(res, repo, "programacion_demo")
        return

    acciones = {
        "1": lambda: menu_crud("PERSONAL", repo.personal, fmt_persona, form_persona, "nombre",
                               repo, Persona, "personal"),
        "2": lambda: menu_crud("VEHÍCULOS", repo.vehiculos, fmt_vehiculo, form_vehiculo, "placa",
                               repo, Vehiculo, "vehiculos"),
        "3": lambda: menu_crud("EQUIPOS", repo.equipos, fmt_equipo,
                               lambda e=None: form_equipo(e, repo.tipos_equipo()), "codigo",
                               repo, Equipo, "equipos"),
        "4": lambda: menu_crud("PUNTOS DE MONITOREO", repo.puntos, fmt_punto,
                               lambda p=None: form_punto(p, repo.tipos_equipo()), "codigo",
                               repo, Punto, "puntos"),
        "5": lambda: menu_config(repo),
        "6": lambda: matriz_tiempos(repo),
        "7": lambda: alertas(repo),
        "8": lambda: menu_planificar(repo),
        "9": menu_historial,
        "10": lambda: cargar_demo(repo) if pedir(
            "Esto reemplaza los datos actuales. ¿Continuar? (s/n)", bool, False) else None,
    }
    while True:
        print(f"""
╔══════════════════════════════════════════════╗
║   PLANIFICADOR DE MONITOREOS AMBIENTALES     ║
╚══════════════════════════════════════════════╝
 Personal: {len(repo.personal)} | Vehículos: {len(repo.vehiculos)} | Equipos: {len(repo.equipos)} | Puntos: {len(repo.puntos)}
  1) Personal
  2) Vehículos
  3) Equipos e instrumentos
  4) Puntos de monitoreo
  5) Configuración (base, jornada, OSRM, costos)
  6) Matriz de tiempos de desplazamiento
  7) Alertas y vencimientos
  8) Generar programación
  9) Historial de programaciones
 10) Cargar datos de ejemplo
  0) Salir""")
        op = input("> ").strip()
        if op == "0":
            break
        if op in acciones:
            try:
                acciones[op]()
            except (KeyboardInterrupt, EOFError):
                print("\n  Operación cancelada.")


if __name__ == "__main__":
    try:
        main()
    except (KeyboardInterrupt, EOFError):
        print("\nHasta luego.")
