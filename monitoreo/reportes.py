"""Salidas: consola, CSV, Excel, hojas de ruta HTML, mapa (folium) y calendario iCalendar (.ics)."""

from __future__ import annotations

import csv
import html
import re
from datetime import date, datetime, timezone
from pathlib import Path

from .distancias import coord_nodo, guardar_cache_osrm, leer_cache_osrm, nombre_nodo, trazado_osrm
from .modelo import (BASE, CARPETA_SALIDAS, DIAS_SEMANA, ESTADOS_TAREA, Configuracion, Repositorio, fmt_min)
from .planificador import Plan, PlanDia, cumplimiento, resumen

COLORES = ["#2563eb", "#dc2626", "#16a34a", "#9333ea", "#ea580c", "#0891b2", "#be185d", "#4d7c0f"]


def nombre_archivo(texto: str) -> str:
    return re.sub(r"[^\w\-]+", "_", texto, flags=re.UNICODE).strip("_")[:60] or "programacion"


def paradas(pd: PlanDia, repo: Repositorio) -> list[tuple[float, float]]:
    return [c for c in (coord_nodo(n, repo) for n in pd.sim.nodos) if c]


# --------------------------------------------------------------------------- #
# Consola
# --------------------------------------------------------------------------- #
def imprimir_plan(plan: Plan, cfg: Configuracion) -> None:
    m = cfg.moneda
    if not plan.dias:
        print("\nNo se pudo programar ningún punto.")
    dia_actual = None
    for pd in plan.dias:
        if pd.dia != dia_actual:
            dia_actual = pd.dia
            print("\n" + "=" * 100)
            print(f" {DIAS_SEMANA[pd.dia.weekday()]} {pd.dia:%d/%m/%Y}")
            print("=" * 100)
        c = pd.cuadrilla
        costo = pd.costos(cfg)
        viaje = f"  [salida {pd.viaje}]" if pd.viaje else ""
        print(f"\n  {c.describir()}{viaje}")
        if c.equipos:
            print(f"  Equipos: {', '.join(c.equipos)}")
        cadena = " → ".join([nombre_nodo(pd.origen, cfg) if pd.origen != BASE else "BASE",
                             *(t.codigo for t in pd.ruta),
                             nombre_nodo(pd.destino, cfg) if pd.destino != BASE else "BASE"])
        print(f"  Ruta: {cadena}   | {pd.km:.1f} km | ~{pd.litros:.1f} L | costo {m} {costo['total']:,.0f}")
        for e in pd.eventos:
            print(f"    {e.inicio:%H:%M}-{e.fin:%H:%M}  {e.actividad:<52} {e.lugar}")
        for a in pd.avisos:
            print(f"    ⚠ {a}")

    r = resumen(plan, cfg)
    print("\n" + "-" * 100 + "\n RESUMEN\n" + "-" * 100)
    print(f"  Tiempos: {plan.fuente_tiempos}")
    for a in plan.avisos_red:
        print(f"  ⚠ {a}")
    print(f"  Puntos: {r['puntos']} (+{r['retiros']} retiros) | Días de campo: {r['dias']} | Noches fuera: "
          f"{r['noches']} | Km: {r['km']:.1f} | Combustible: {r['litros']:.1f} L")
    print(f"  Costos: combustible {m} {r['combustible']:,.0f} + vehículos {m} {r['vehiculo']:,.0f} + viáticos "
          f"{m} {r['viaticos']:,.0f} + alojamiento {m} {r['alojamiento']:,.0f} = TOTAL {m} {r['total']:,.0f}")
    print("  Por persona: " + "; ".join(f"{n} {d['dias']}d/{fmt_min(d['min'])}"
                                         for n, d in sorted(r["personas"].items())))
    print("  Por vehículo: " + "; ".join(f"{p} {d['dias']}d/{d['km']:.0f} km"
                                          for p, d in sorted(r["flota"].items())))
    if plan.pendientes:
        print("\n  NO PROGRAMADOS:")
        for t in plan.pendientes:
            print(f"   ✗ {t.codigo} ({t.tipo}): {plan.motivos.get(t.clave, '')}")


# --------------------------------------------------------------------------- #
# CSV y Excel
# --------------------------------------------------------------------------- #
def exportar_csv(plan: Plan, ruta: Path) -> None:
    with ruta.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Fecha", "Cuadrilla", "Vehículo", "Conductor", "Integrantes", "Equipos",
                    "Inicio", "Fin", "Duración (min)", "Actividad", "Lugar", "Km", "Estado", "Hora real",
                    "Observación"])
        for pd in plan.dias:
            c = pd.cuadrilla
            for e in pd.eventos:
                t = e.tarea
                w.writerow([pd.dia.isoformat(), c.id, c.vehiculo.placa, c.conductor.nombre if c.conductor else "",
                            ", ".join(x.nombre for x in c.integrantes), ", ".join(c.equipos),
                            f"{e.inicio:%H:%M}", f"{e.fin:%H:%M}", round((e.fin - e.inicio).total_seconds() / 60),
                            e.actividad, e.lugar, f"{e.km:.1f}".replace(".", ","),
                            ESTADOS_TAREA.get(t.estado, "") if t else "", t.hora_real if t else "",
                            t.observacion if t else ""])


def exportar_excel(plan: Plan, cfg: Configuracion, ruta: Path) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return False
    wb = Workbook()
    wb.remove(wb.active)
    cab_font = Font(bold=True, color="FFFFFF")
    cab_fill = PatternFill("solid", fgColor="1F6F5C")
    aviso_fill = PatternFill("solid", fgColor="FDECEA")
    m = cfg.moneda

    def hoja(titulo: str, encabezados: list[str], filas: list[list]):
        ws = wb.create_sheet(titulo)
        ws.append(encabezados)
        for celda in ws[1]:
            celda.font, celda.fill = cab_font, cab_fill
            celda.alignment = Alignment(vertical="center", wrap_text=True)
        for f in filas:
            ws.append(f)
        ws.freeze_panes = "A2"
        if filas:
            ws.auto_filter.ref = ws.dimensions
        for i, col in enumerate(ws.columns, 1):
            largo = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, largo + 2), 60)
        for fila in ws.iter_rows(min_row=2):
            for c in fila:
                if isinstance(c.value, date):
                    c.number_format = "dd/mm/yyyy"
                elif isinstance(c.value, float):
                    c.number_format = "#,##0.00"
        return ws

    r = resumen(plan, cfg)
    cu = cumplimiento(plan)
    hoja("Resumen", ["Indicador", "Valor"], [
        ["Programación", plan.nombre], ["Estado", plan.estado], ["Fecha de inicio", plan.fecha_inicio],
        ["Fuente de tiempos", plan.fuente_tiempos], ["Puntos programados", r["puntos"]],
        ["Retiros de equipos", r["retiros"]], ["Días de campo", r["dias"]], ["Noches fuera", r["noches"]],
        ["Km totales", round(r["km"], 1)], ["Combustible (L)", round(r["litros"], 1)],
        [f"Costo combustible ({m})", round(r["combustible"], 2)], [f"Costo vehículos ({m})", round(r["vehiculo"], 2)],
        [f"Viáticos ({m})", round(r["viaticos"], 2)], [f"Alojamiento ({m})", round(r["alojamiento"], 2)],
        [f"COSTO TOTAL ({m})", round(r["total"], 2)], ["Tareas no programadas", r["pendientes"]],
        ["Avisos de validación", r["avisos"]], ["Cumplimiento (% ejecutado)", cu["porcentaje"]],
    ])
    filas = []
    for pd in plan.dias:
        c, costo = pd.cuadrilla, pd.costos(cfg)
        filas.append([pd.dia, DIAS_SEMANA[pd.dia.weekday()], c.id, pd.viaje, c.vehiculo.placa,
                      c.conductor.nombre if c.conductor else "", ", ".join(x.nombre for x in c.integrantes),
                      ", ".join(c.equipos), nombre_nodo(pd.origen, cfg),
                      " → ".join(f"{t.codigo}{'(R)' if t.tipo == 'retiro' else '(I)' if t.tipo == 'instalacion' else ''}"
                                 for t in pd.ruta),
                      nombre_nodo(pd.destino, cfg),
                      f"{pd.eventos[0].inicio:%H:%M}" if pd.eventos else "",
                      f"{pd.eventos[-1].fin:%H:%M}" if pd.eventos else "", round(pd.km, 1), round(pd.litros, 1),
                      round(costo["combustible"], 2), round(costo["vehiculo"], 2), round(costo["viaticos"], 2),
                      round(costo["alojamiento"], 2), round(costo["total"], 2), " | ".join(pd.avisos)])
    ws = hoja("Rutas", ["Fecha", "Día", "Cuadrilla", "Salida", "Vehículo", "Conductor", "Integrantes", "Equipos",
                        "Desde", "Secuencia (I=instalación, R=retiro)", "Hasta", "Inicio", "Fin", "Km", "Litros",
                        f"Combustible {m}", f"Vehículo {m}", f"Viáticos {m}", f"Alojamiento {m}", f"Total {m}",
                        "Avisos"], filas)
    for fila in ws.iter_rows(min_row=2):
        if fila[-1].value:
            for c in fila:
                c.fill = aviso_fill
    hoja("Cronograma", ["Fecha", "Cuadrilla", "Vehículo", "Inicio", "Fin", "Min", "Actividad", "Lugar", "Km",
                        "Latitud", "Longitud"],
         [[pd.dia, pd.cuadrilla.id, pd.cuadrilla.vehiculo.placa, f"{e.inicio:%H:%M}", f"{e.fin:%H:%M}",
           round((e.fin - e.inicio).total_seconds() / 60), e.actividad, e.lugar, round(e.km, 1),
           e.tarea.punto.lat if e.tarea else None, e.tarea.punto.lon if e.tarea else None]
          for pd in plan.dias for e in pd.eventos])
    horas = {id(e.tarea): e for pd in plan.dias for e in pd.eventos if e.tarea}
    hoja("Registro de campo", ["Fecha", "Cuadrilla", "Punto", "Nombre", "Matriz", "Actividad", "Hora programada",
                               "Estado", "Hora real", "Observación"],
         [[pd.dia, pd.cuadrilla.id, t.codigo, t.punto.nombre, t.punto.matriz, t.etiqueta(),
           f"{horas[id(t)].inicio:%H:%M}" if id(t) in horas else "", ESTADOS_TAREA.get(t.estado, t.estado),
           t.hora_real, t.observacion] for pd in plan.dias for t in pd.ruta])
    hoja("Personal", ["Persona", "Días de campo", "Horas", "Noches fuera"],
         [[n, d["dias"], round(d["min"] / 60, 1), d["noches"]] for n, d in sorted(r["personas"].items())])
    hoja("Vehículos", ["Placa", "Días", "Km", "Litros", f"Costo {m} (comb.+vehículo)"],
         [[p, d["dias"], round(d["km"], 1), round(d["litros"], 1), round(d["costo"], 2)]
          for p, d in sorted(r["flota"].items())])
    hoja("No programados", ["Punto", "Tarea", "Motivo"],
         [[t.codigo, t.tipo, plan.motivos.get(t.clave, "")] for t in plan.pendientes])
    wb.save(str(ruta))
    return True


# --------------------------------------------------------------------------- #
# Hojas de ruta (HTML imprimible)
# --------------------------------------------------------------------------- #
def exportar_hoja_ruta(plan: Plan, repo: Repositorio, ruta: Path) -> None:
    cfg, e_ = repo.config, html.escape
    secciones = []
    for pd in plan.dias:
        c, costo = pd.cuadrilla, pd.costos(cfg)
        ruta_gmaps = "https://www.google.com/maps/dir/" + "/".join(f"{a},{b}" for a, b in paradas(pd, repo))
        filas = []
        for ev in pd.eventos:
            nota, enlace = "", ""
            if ev.tarea:
                p = ev.tarea.punto
                detalles = [f"{p.lat:.5f}, {p.lon:.5f}"]
                if p.hora_desde or p.hora_hasta:
                    detalles.append(f"ventana {p.hora_desde or '--'}–{p.hora_hasta or '--'}")
                if p.max_horas_preservacion and ev.tarea.recoge_muestra:
                    detalles.append(f"preservar ≤ {p.max_horas_preservacion:g} h")
                if p.entrega_laboratorio and cfg.usar_laboratorio and ev.tarea.recoge_muestra:
                    detalles.append("entregar al laboratorio")
                if p.equipos and ev.tarea.tipo != "retiro":
                    detalles.append("equipos: " + ", ".join(p.equipos))
                if p.requiere_4x4:
                    detalles.append("acceso 4x4")
                nota = f"<div class='nota'>{e_(' · '.join(detalles))}</div>"
                enlace = (f"<a href='https://www.google.com/maps/dir/?api=1&destination={p.lat},{p.lon}"
                          f"&travelmode=driving' target='_blank'>Navegar</a>")
            filas.append(f"<tr class='{'muestra' if ev.tarea else ''}'><td>{ev.inicio:%H:%M}</td>"
                         f"<td>{ev.fin:%H:%M}</td><td>{e_(ev.actividad)}{nota}</td><td>{e_(ev.lugar)}</td>"
                         f"<td>{enlace}</td><td class='chk'></td></tr>")
        avisos = "".join(f"<li>{e_(a)}</li>" for a in pd.avisos)
        noche = (f"<dt>Pernocte</dt><dd>{e_(nombre_nodo(pd.destino, cfg))} (salida {e_(pd.viaje)})</dd>"
                 if pd.pernocta else "")
        secciones.append(f"""
<section>
  <header>
    <div><h2>{DIAS_SEMANA[pd.dia.weekday()]} {pd.dia:%d/%m/%Y} · Cuadrilla {e_(c.id)}</h2>
    <p class="sub">Sale de {e_(nombre_nodo(pd.origen, cfg))} · {pd.eventos[0].inicio:%H:%M} → fin {pd.eventos[-1].fin:%H:%M} en {e_(nombre_nodo(pd.destino, cfg))}</p></div>
    <a class="btn" href="{ruta_gmaps}" target="_blank">Ruta completa en Google Maps</a>
  </header>
  {f'<ul class="avisos">{avisos}</ul>' if avisos else ''}
  <dl>
    <dt>Vehículo</dt><dd>{e_(c.vehiculo.placa)} · {e_(c.vehiculo.descripcion)}</dd>
    <dt>Conductor</dt><dd>{e_(c.conductor.nombre if c.conductor else '—')}</dd>
    <dt>Integrantes</dt><dd>{e_(', '.join(f"{x.nombre}{' · ' + x.telefono if x.telefono else ''}" for x in c.integrantes))}</dd>
    <dt>Equipos</dt><dd>{e_(', '.join(c.equipos) or '—')}</dd>{noche}
    <dt>Recorrido</dt><dd>{pd.km:.1f} km · ~{pd.litros:.1f} L · {cfg.moneda} {costo['total']:,.0f}</dd>
  </dl>
  <table>
    <thead><tr><th>Inicio</th><th>Fin</th><th>Actividad</th><th>Lugar</th><th></th><th>✓</th></tr></thead>
    <tbody>{''.join(filas)}</tbody>
  </table>
  <div class="firmas"><div>Responsable de cuadrilla</div><div>Supervisor</div></div>
</section>""")
    doc = f"""<!doctype html>
<html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Hojas de ruta · {e_(plan.nombre)}</title>
<style>
  body {{ font: 14px/1.4 system-ui, Segoe UI, sans-serif; color: #1c2b27; background: #f4f6f5; margin: 0; padding: 16px; }}
  section {{ background: #fff; max-width: 960px; margin: 0 auto 24px; padding: 20px 24px; border-radius: 8px;
            box-shadow: 0 1px 3px rgba(0,0,0,.12); }}
  header {{ display: flex; justify-content: space-between; align-items: flex-start; gap: 12px; flex-wrap: wrap; }}
  h1 {{ max-width: 960px; margin: 0 auto 16px; font-size: 20px; }}
  h2 {{ margin: 0; font-size: 18px; color: #1f6f5c; }}
  .sub {{ margin: 2px 0 0; color: #5a6b66; }}
  .btn {{ background: #1f6f5c; color: #fff; padding: 6px 12px; border-radius: 6px; text-decoration: none; font-size: 13px; }}
  .avisos {{ background: #fdecea; color: #8a1c12; border-radius: 6px; padding: 8px 8px 8px 28px; margin: 12px 0 0; }}
  dl {{ display: grid; grid-template-columns: 110px 1fr; gap: 2px 12px; margin: 14px 0; }}
  dt {{ font-weight: 600; color: #5a6b66; }} dd {{ margin: 0; }}
  table {{ width: 100%; border-collapse: collapse; }}
  th, td {{ border-bottom: 1px solid #dde3e1; padding: 5px 6px; text-align: left; vertical-align: top; }}
  th {{ background: #eef3f1; font-size: 12px; text-transform: uppercase; color: #5a6b66; }}
  tr.muestra td {{ background: #f0f8f5; font-weight: 600; }}
  .nota {{ font-weight: 400; font-size: 12px; color: #5a6b66; }}
  td.chk {{ width: 28px; border-left: 1px solid #dde3e1; }}
  a {{ color: #1f6f5c; }}
  .firmas {{ display: flex; gap: 40px; margin-top: 48px; }}
  .firmas div {{ flex: 1; border-top: 1px solid #1c2b27; padding-top: 4px; text-align: center; color: #5a6b66; }}
  @media print {{ body {{ background: #fff; padding: 0; }} h1, .btn {{ display: none; }}
                  section {{ box-shadow: none; page-break-after: always; margin: 0; max-width: none; }} }}
  @media (max-width: 600px) {{ section {{ padding: 14px; }} dl {{ grid-template-columns: 1fr; }} }}
</style></head><body>
<h1>Hojas de ruta · {e_(plan.nombre)}</h1>
{''.join(secciones) or '<p>Sin rutas programadas.</p>'}
</body></html>"""
    ruta.write_text(doc, encoding="utf-8")


# --------------------------------------------------------------------------- #
# Mapa
# --------------------------------------------------------------------------- #
def trazados(plan: Plan, repo: Repositorio) -> list[dict]:
    """Trazado por carretera de cada día (o línea recta si no hay OSRM)."""
    cfg = repo.config
    cache = leer_cache_osrm() if cfg.usar_osrm else {}
    salida = []
    for i, pd in enumerate(plan.dias):
        coords = paradas(pd, repo)
        linea = trazado_osrm(coords, cfg.osrm_url, cache) if cfg.usar_osrm else None
        salida.append({"indice": i, "color": COLORES[i % len(COLORES)], "linea": linea or [list(c) for c in coords],
                       "recta": linea is None})
    if cfg.usar_osrm:
        guardar_cache_osrm(cache)
    return salida


def exportar_mapa(plan: Plan, repo: Repositorio, ruta: Path) -> bool:
    try:
        import folium
    except ImportError:
        return False
    cfg = repo.config
    m = folium.Map(location=[cfg.base_lat, cfg.base_lon], zoom_start=10)
    folium.Marker([cfg.base_lat, cfg.base_lon], tooltip=cfg.base_nombre,
                  icon=folium.Icon(color="black", icon="home")).add_to(m)
    if cfg.usar_laboratorio and (cfg.lab_lat or cfg.lab_lon):
        folium.Marker([cfg.lab_lat, cfg.lab_lon], tooltip=cfg.lab_nombre,
                      icon=folium.Icon(color="darkblue", icon="flask", prefix="fa")).add_to(m)
    for a in repo.alojamientos:
        if a.activo:
            folium.Marker([a.lat, a.lon], tooltip=f"Alojamiento: {a.nombre}",
                          icon=folium.Icon(color="gray", icon="bed", prefix="fa")).add_to(m)
    rectas = 0
    for tz, pd in zip(trazados(plan, repo), plan.dias):
        col = tz["color"]
        capa = folium.FeatureGroup(name=f"{pd.dia:%d/%m} {pd.cuadrilla.id} ({pd.km:.0f} km)")
        if tz["recta"]:
            rectas += 1
        folium.PolyLine(tz["linea"], color=col, weight=4 if not tz["recta"] else 3, opacity=0.8,
                        dash_array="8" if tz["recta"] else None,
                        tooltip=f"{pd.dia:%d/%m} {pd.cuadrilla.id}").add_to(capa)
        horas = {id(e.tarea): e for e in pd.eventos if e.tarea}
        for n, t in enumerate(pd.ruta, 1):
            p, e = t.punto, horas.get(id(t))
            hora = f" | {e.inicio:%H:%M}-{e.fin:%H:%M}" if e else ""
            folium.CircleMarker([p.lat, p.lon], radius=8, color=col, fill=True, fill_opacity=0.9,
                                tooltip=f"{n}. {p.codigo} - {p.nombre} · {t.etiqueta()}{hora}").add_to(capa)
            folium.Marker([p.lat, p.lon], icon=folium.DivIcon(
                html=f'<div style="font:bold 10px sans-serif;color:#fff;text-align:center;'
                     f'width:16px;margin:-6px 0 0 -8px">{n}</div>')).add_to(capa)
        capa.add_to(m)
    if rectas:
        print(f"  ⚠ {rectas} ruta(s) sin trazado OSRM en el mapa; se dibujan con línea recta punteada.")
    folium.LayerControl(collapsed=False).add_to(m)
    m.save(str(ruta))
    return True


# --------------------------------------------------------------------------- #
# Calendario (iCalendar .ics: Google Calendar, Outlook, Apple)
# --------------------------------------------------------------------------- #
def _ics_texto(s: str) -> str:
    return s.replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,").replace("\n", "\\n")


def _ics_plegar(linea: str) -> str:
    """Pliega líneas a 75 octetos como exige RFC 5545."""
    salida, actual = [], b""
    for ch in linea:
        b = ch.encode("utf-8")
        if len(actual) + len(b) > (75 if not salida else 74):
            salida.append(actual.decode("utf-8"))
            actual = b""
        actual += b
    salida.append(actual.decode("utf-8"))
    return "\r\n ".join(salida)


def exportar_ics(plan: Plan, repo: Repositorio, ruta: Path, persona: str | None = None) -> int:
    """Un evento por cuadrilla y día. Con `persona`, solo los días en que participa.
    Las horas se escriben en hora local (sin zona), que es como las interpreta el calendario al importar."""
    cfg = repo.config
    ahora = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    lineas = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//Planificador de Monitoreos Ambientales//ES",
              "CALSCALE:GREGORIAN", "METHOD:PUBLISH", f"X-WR-CALNAME:{_ics_texto(plan.nombre)}"]
    n = 0
    for pd in plan.dias:
        c = pd.cuadrilla
        if persona and persona not in {x.nombre for x in c.integrantes}:
            continue
        if not pd.eventos:
            continue
        n += 1
        puntos = ", ".join(t.codigo for t in pd.ruta) or f"retorno desde {nombre_nodo(pd.origen, cfg)}"
        detalle = [f"Cuadrilla {c.id} · vehículo {c.vehiculo.placa}",
                   f"Conductor: {c.conductor.nombre if c.conductor else '—'}",
                   f"Integrantes: {', '.join(x.nombre for x in c.integrantes)}",
                   f"Equipos: {', '.join(c.equipos) or '—'}", ""]
        detalle += [f"{e.inicio:%H:%M}-{e.fin:%H:%M} {e.actividad} · {e.lugar}" for e in pd.eventos]
        if pd.pernocta:
            detalle += ["", f"PERNOCTE en {nombre_nodo(pd.destino, cfg)}"]
        if pd.avisos:
            detalle += ["", "AVISOS: " + "; ".join(pd.avisos)]
        detalle += ["", "Ruta: https://www.google.com/maps/dir/" + "/".join(f"{a},{b}" for a, b in paradas(pd, repo))]
        uid = f"{plan.id}-{pd.dia:%Y%m%d}-{c.id}@planificador-monitoreo"
        ev = ["BEGIN:VEVENT", f"UID:{uid}", f"DTSTAMP:{ahora}",
              f"DTSTART:{pd.eventos[0].inicio:%Y%m%dT%H%M%S}", f"DTEND:{pd.eventos[-1].fin:%Y%m%dT%H%M%S}",
              f"SUMMARY:{_ics_texto(f'Monitoreo {c.id}: {puntos}')}",
              f"LOCATION:{_ics_texto(nombre_nodo(pd.origen, cfg))}",
              f"DESCRIPTION:{_ics_texto(chr(10).join(detalle))}", "STATUS:CONFIRMED" if plan.estado != "borrador"
              else "STATUS:TENTATIVE"]
        for x in c.integrantes:
            if x.email:
                ev.append(f"ATTENDEE;CN={_ics_texto(x.nombre)};ROLE=REQ-PARTICIPANT:mailto:{x.email}")
        ev += ["BEGIN:VALARM", "ACTION:DISPLAY", "DESCRIPTION:Salida a campo", "TRIGGER:-PT12H", "END:VALARM",
               "END:VEVENT"]
        lineas += ev
    lineas.append("END:VCALENDAR")
    ruta.write_text("\r\n".join(_ics_plegar(x) for x in lineas) + "\r\n", encoding="utf-8", newline="")
    return n


# --------------------------------------------------------------------------- #
# Exportación conjunta
# --------------------------------------------------------------------------- #
def exportar(plan: Plan, repo: Repositorio, formato: str, persona: str | None = None) -> Path:
    """Genera un archivo en la carpeta de salidas y devuelve su ruta."""
    CARPETA_SALIDAS.mkdir(exist_ok=True)
    base = CARPETA_SALIDAS / nombre_archivo(plan.nombre)
    if formato == "csv":
        destino = base.with_suffix(".csv")
        exportar_csv(plan, destino)
    elif formato == "xlsx":
        destino = base.with_suffix(".xlsx")
        if not exportar_excel(plan, repo.config, destino):
            raise RuntimeError("Falta el paquete 'openpyxl' para exportar a Excel")
    elif formato == "hojas":
        destino = base.with_name(base.name + "_hojas_de_ruta.html")
        exportar_hoja_ruta(plan, repo, destino)
    elif formato == "mapa":
        destino = base.with_name(base.name + "_mapa.html")
        if not exportar_mapa(plan, repo, destino):
            raise RuntimeError("Falta el paquete 'folium' para generar el mapa")
    elif formato == "ics":
        sufijo = f"_{nombre_archivo(persona)}" if persona else ""
        destino = base.with_name(base.name + sufijo + ".ics")
        exportar_ics(plan, repo, destino, persona)
    else:
        raise ValueError(f"formato desconocido: {formato}")
    return destino


def exportar_todo(plan: Plan, repo: Repositorio) -> list[Path]:
    rutas = []
    for fmt in ("csv", "xlsx", "hojas", "mapa", "ics"):
        try:
            rutas.append(exportar(plan, repo, fmt))
        except RuntimeError as e:
            print(f"  ({e})")
    personas = sorted({x.nombre for pd in plan.dias for x in pd.cuadrilla.integrantes})
    for p in personas:
        rutas.append(exportar(plan, repo, "ics", p))
    return rutas
